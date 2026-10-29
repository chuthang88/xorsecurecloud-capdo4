from pathlib import Path
from docx import Document
import re,json
root=Path('output/XOR-Ho-so-CD4-v2.4');out=Path('tmp/ho-so-v2-4')
texts={}
for p in root.rglob('*.docx'):
 d=Document(p)
 txt='\n'.join([x.text for x in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
 texts[str(p)]=txt
 if 'Bo-2' in str(p):
  for line in txt.splitlines():
   if re.search('dự thảo|mẫu|hướng dẫn|không tự|nguồn cung cấp|workbook|OpenStack|trước khi trình|bản này|bản nguồn|được cấp phép|chưa ký|điền kết quả thật',line,re.I):print(p.name,':',line)
(out/'all-doc-text.json').write_text(json.dumps(texts,ensure_ascii=False,indent=2))
print('DOCX',len(texts))
