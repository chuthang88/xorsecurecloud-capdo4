from pathlib import Path
import json, hashlib, re, shutil, zipfile
from collections import Counter
from docx import Document
from openpyxl import load_workbook
from inventory_data import T,ASSETS,SHA
P=Path(__file__).resolve().parent
ROOT=P.parents[1]; OLD=ROOT/'output/XOR-Ho-so-CD4-v2.4'; OUT=ROOT/'output/XOR-Ho-so-CD4-v2.5'
report={};errors=[]
def check(label,ok,detail=None):
 report[label]={'pass':bool(ok),'detail':detail}
 if not ok:errors.append(label)
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
check('source_unchanged',sha('/Users/thangchu/Work/XOR/Xor Secure Cloud/Bao-cao-kiem-ke-ha-tang-he-thong-XOR-Secure-20261018.docx')==SHA)
changed=set(json.loads((P/'changed-docs.json').read_text()))
manifest=list(map(Path,json.loads((P/'doc_manifest.json').read_text())))
fulltexts={}
for p in manifest:
 d=Document(p);text='\n'.join(x.text for x in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
 fulltexts[p.name]=text
 check('brand_'+p.name,'OpenStack' not in text and 'dựa trên OpenStack' not in text)
 check('version_'+p.name,'Phiên bản: 2.4' not in text and 'phiên bản 2.4' not in text)
 if str(p) not in changed:check('preserved_'+p.name,sha(p)==sha(OLD/p.relative_to(OUT)))
name='02_Phu_luc_doi_chieu_yeu_cau_TCVN.docx'
def rows(path):
 return {r.cells[0].text.split('\n')[0]:[c.text for c in r.cells] for t in Document(path).tables for r in t.rows if r.cells[0].text.startswith('TC-')}
a=rows(OLD/'Bo-2-Hanh-chinh'/name);b=rows(OUT/'Bo-2-Hanh-chinh'/name)
check('119_requirements_preserved',len(a)==len(b)==119 and all(a[k][:2]==b[k][:2] and a[k][2].split('Hiện trạng:')[0]==b[k][2].split('Hiện trạng:')[0] for k in a))
check('119_word_verdicts_pending',all('Kết quả: Chưa đủ bằng chứng.' in r[3] for r in b.values()))
check('main_current_state_consistent',all(x in fulltexts['01_Thuyet_minh_ho_so_de_xuat_cap_do_4.docx'] for x in ['2.336 GiB','46 VM','03 controller','EV-001','19,20 TB']))
ann=fulltexts['10_Phu_luc_thiet_ke_va_hien_trang.docx']
serials=[r[2] for r in T[1][1:]]+[r[3] for r in T[4][1:]]
check('13_serials_in_annex',len(set(serials))==13 and all(x in ann for x in serials))
vmnames=[r[1] for t in [13,14,15] for r in T[t][1:]]+[r[0] for t in [16,17] for r in T[t][1:]]+['mfa-02']
check('46_vms_in_annex',len(vmnames)==46 and all(n in ann for n in vmnames))
check('inventory_totals',sum(int(r[4]) for r in T[1][1:])==2336 and sum(int(r[2]) for t in [13,14,15] for r in T[t][1:])==166 and sum(float(r[3]) for t in [13,14,15] for r in T[t][1:])==340 and sum(int(r[3]) for t in [16,17] for r in T[t][1:])==113 and sum(float(r[4]) for t in [16,17] for r in T[t][1:])==215.5)
check('dr_limitation_explicit','quan hệ với DR Nam Thăng Long' in ann and 'CẦN XÁC NHẬN nhà cung cấp' in ann)
log=json.loads((P/'xlsx-edit-log.json').read_text());allowed={(r[0],r[1]) for r in log}
legacyrows=json.loads((P/'legacy-edited-rows.json').read_text())
for file in ['02_Checklist_ra_soat_XOR_cap_do_4.xlsx','04_Doi_soat_129_hang_muc_v094.xlsx']:
 x=load_workbook(OLD/'Bo-1-Ky-thuat'/file);y=load_workbook(OUT/'Bo-1-Ky-thuat'/file);yc=load_workbook(OUT/'Bo-1-Ky-thuat'/file,data_only=True)
 def formulas(w):return {(s.title,c.coordinate):c.value for s in w for row in s for c in row if c.data_type=='f'}
 fx,fy=formulas(x),formulas(y)
 check(file+'_formulas',fx==fy,{'before':len(fx),'after':len(fy)})
 check(file+'_sheets',x.sheetnames==y.sheetnames,y.sheetnames)
 bad=[];changes=[]
 for s in y:
  old=x[s.title]
  for row in s:
   for c in row:
    v0=old[c.coordinate].value;v1=c.value
    if v0==v1 or (v0 in [None,''] and v1 in [None,'']):continue
    changes.append((s.title,c.coordinate))
    if file.startswith('02'):
     okay=(s.title,c.coordinate) in allowed or (s.title=='Tai san' and 6<=c.row<=64 and c.column<=15) or (s.title=='Bang chung' and 6<=c.row<=12 and 2<=c.column<=12)
    else:okay=s.title=='Doi_soat' and c.row in legacyrows and c.column in [9,12,13]
    if not okay:bad.append([s.title,c.coordinate,str(v0),str(v1)])
 check(file+'_only_intended_value_changes',not bad,{'changes':len(changes),'unintended':bad})
 errs=[(s.title,c.coordinate,c.value) for s in yc for row in s for c in row if c.data_type=='e']
 check(file+'_no_cached_errors',not errs,errs)
 validation=[]
 for s in y:
  if s.title=='Tai san' and file.startswith('02'):continue
  if str(s.data_validations)!=str(x[s.title].data_validations):validation.append(s.title)
 check(file+'_validations_preserved',not validation,validation)
 if file.startswith('02'):
  values=list(y['Tai san'].iter_rows(min_row=6,max_row=64,max_col=15,values_only=True))
  expected=[tuple(r[:13]+['Cần xác minh']+r[14:]) for r in ASSETS]
  check('59_asset_rows_match',values==expected and len({r[0] for r in values})==59)
  check('asset_table_expanded',list(y['Tai san'].tables.values())[0].ref=='A5:O64',list(y['Tai san'].tables.values())[0].ref)
  check('asset_validation_expanded',any(str(d.sqref)=='N6:N64' for d in y['Tai san'].data_validations.dataValidation))
  check('evidence_receipt_only',y['Bang chung']['L6'].value=='Đã thu thập' and all(y['Bang chung'][f'L{r}'].value=='Chưa thu thập' for r in range(7,13)))
  check('tcvn_verdicts_preserved',all(y['TCVN'].cell(r,c).value==x['TCVN'].cell(r,c).value for r in range(6,125) for c in [10,12,13,14]))
 else:
  check('historical_sheet_unchanged',all(c.value==x['Nguon_v094'][c.coordinate].value and c._style==x['Nguon_v094'][c.coordinate]._style for row in y['Nguon_v094'] for c in row))
  check('legacy_status_unchanged',all(y['Doi_soat'].cell(r,c).value==x['Doi_soat'].cell(r,c).value for r in range(1,y['Doi_soat'].max_row+1) for c in [6,10,11]))
results=json.loads((P/'render-results.json').read_text());oldresults=json.loads((ROOT/'tmp/ho-so-v2-4/render-results.json').read_text())
pagecount=sum(r['pages'] for r in results)+sum(r['pages'] for r in oldresults if r['index'] not in {v['index'] for v in results})
check('rendered_changed_docs',len(results)==7 and all(r['code']==0 for r in results),{'total_pages':pagecount,'changed_doc_pages':sum(r['pages'] for r in results)})
(P/'final-qa.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps({'checks':len(report),'failed':errors,'page_count':pagecount},ensure_ascii=False))
if errors:raise SystemExit(1)
