from doc_base import *
from inventory_data import *
from docx.enum.section import WD_SECTION_START
from docx.text.paragraph import Paragraph

OLD=ROOT/'output/XOR-Ho-so-CD4-v2.4'
changed=[]
def writep(p,text):
 rpr=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 p.clear();r=p.add_run(text)
 if rpr is not None:r._r.insert(0,rpr)
def cell(c,text):
 p=c.paragraphs[0];writep(p,text)
 for q in c.paragraphs[1:]:q._p.getparent().remove(q._p)
def replace(d,old,new):
 n=0
 for p in d.paragraphs:
  if old in p.text:writep(p,p.text.replace(old,new));n+=1
 for t in d.tables:
  for row in t.rows:
   for c in row.cells:
    for p in c.paragraphs:
     if old in p.text:writep(p,p.text.replace(old,new));n+=1
 return n
def save(d,p):
 replace(d,'phiên bản 2.4','phiên bản 2.5')
 replace(d,'Phiên bản: 2.4.','Phiên bản: 2.5.')
 d.save(p);changed.append(str(p))
def rowmap(t,m):
 for r in t.rows[1:]:
  if r.cells[0].text in m:cell(r.cells[1],m[r.cells[0].text])

# Main submission dossier: populate the existing current-state fields.
p=ADMIN/'01_Thuyet_minh_ho_so_de_xuat_cap_do_4.docx';d=Document(p)
rowmap(d.tables[2],{
 'Quy mô hai IDC':SUMMARY+' Ngoài site mới xác định 01 VM; cấu hình DR đầy đủ, rack và tải thực: [CẦN ĐIỀN].',
 'Đối tượng phục vụ':'Có 12 project trong danh mục định danh, trong đó 07 project có VM tại thời điểm kiểm kê. [CẦN ĐIỀN số khách hàng, người dùng, địa bàn và quan hệ project–khách hàng].',
 'Dịch vụ':'Đã ghi nhận VM, mạng/IP, 26 volume với 2.387 GiB logic và 19 image. [CẦN XÁC NHẬN danh mục dịch vụ đã đưa vào cung cấp; trạng thái Kubernetes, GPU, object/file và DBaaS].'})
replace(d,'Hiện trạng triển khai: [CẦN ĐIỀN số, ngày hồ sơ kiểm kê/nghiệm thu; phiên bản cấu hình; sai khác giữa thiết kế và hiện trạng; mã tài liệu tại PL01–PL07].','Hiện trạng được cập nhật theo '+SOURCE_TITLE+' (EV-001); danh mục chi tiết tại tài liệu 10. Hình 1–2 thể hiện phương án thiết kế. Sơ đồ tài sản kiểm kê tại tài liệu 10 phân biệt phần đã ghi nhận và phần còn phải xác nhận. [CẦN ĐIỀN số biên bản nghiệm thu và văn bản phê duyệt thiết kế].')
replace(d,'Hình 1. Kiến trúc tổng thể hệ thống XOR SECURE CLOUD.','Hình 1. Phương án kiến trúc tổng thể hệ thống XOR SECURE CLOUD.')
replace(d,'Hình 2. Phân vùng mạng và các luồng kết nối chính.','Hình 2. Phương án phân vùng mạng và các luồng kết nối chính.')
replace(d,'Mỗi IDC được thiết kế thành cụm độc lập, có quorum nội bộ cho các thành phần yêu cầu đồng thuận.','Phương án mục tiêu bố trí mỗi IDC thành cụm độc lập, có quorum nội bộ cho các thành phần yêu cầu đồng thuận. Hiện trạng kiểm kê ghi nhận 03 controller là VM trên 03 máy chủ Proxmox tại site chính; cấu hình cụm tương ứng tại DR chưa được xác nhận.')
for i,p0 in enumerate(d.paragraphs):
 if p0.text.startswith('Mã bản ghi: [CẦN ĐIỀN]. Ngày chốt hiện trạng:'):
  # Keep untouched date for data classification/flows lacking evidence.
  prev=d.paragraphs[i-1].text
  if any(k in prev for k in ['PL01','PL02','PL03','PL06','PL07']):writep(p0,'Mã bản ghi: EV-001. Thời điểm kiểm kê: 18/10/2026, khoảng 21:19–21:31 GMT+7.')
rowmap(d.tables[11],{
 'Mã dịch vụ; tên; chức năng':'Ghi nhận VM, mạng ảo, image, volume và VM quản trị. [CẦN ĐIỀN mã và danh mục dịch vụ được phê duyệt].',
 'Đối tượng và phạm vi phục vụ':'12 project; 07 project có VM. [CẦN ĐIỀN nhóm khách hàng, địa bàn và số người dùng].',
 'Thành phần và vị trí triển khai':'Site chính: 21 VM Proxmox, gồm 03 controller, và 24 VM compute; ngoài site: mfa-02. Chi tiết tại tài liệu 10.',
 'Năng lực và nghĩa vụ dịch vụ':'VM Proxmox: 166 vCPU, 340 GiB RAM cấp; VM compute: 113 vCPU, 215,5 GiB RAM cấp; 26 volume, 2.387 GiB logic. [CẦN ĐIỀN tải thực, SLA và BIA].',
 'Tài liệu xác nhận':'EV-001, tr.9–15; tài liệu 10. [CẦN ĐIỀN hợp đồng dịch vụ và phê duyệt danh mục].'})
rowmap(d.tables[12],{
 'Mã tài sản; loại; vai trò':'KK-SRV-01–09: 03 Proxmox, 03 compute, 03 Ceph; KK-NET-01–04: 02 Nexus, 02 FortiGate. Chi tiết tại tài liệu 10.',
 'Site và vị trí':'Site chính, đối chiếu DC Cầu Giấy theo phạm vi hồ sơ. [CẦN ĐIỀN phòng, rack/U và biên bản xác nhận vị trí].',
 'Nhà sản xuất; model; serial':'09 Dell PowerEdge R630; 02 Cisco N9K-C93180YC-EX; 02 FortiGate-500E. Đủ 13 serial tại tài liệu 10, mục 2–3.',
 'Cấu hình và phiên bản':'324 nhân/648 luồng; RAM DIMM 2.336 GiB. Nexus: NX-OS 9.3(8); FortiGate: FortiOS 7.6.6 build3652. [CẦN ĐIỀN phiên bản host/hypervisor].',
 'Kết nối vật lý':'Bond LACP hoặc active-backup, thành viên 10/25 Gb theo host; xem tài liệu 10. [CẦN ĐIỀN port map, nhãn cáp và fabric].',
 'Trách nhiệm và hiện trạng':'Kiểm kê ngày 18/10/2026. [CẦN ĐIỀN chủ sở hữu, người quản lý và trạng thái vận hành từng thiết bị].',
 'Hồ sơ hiện trạng':'EV-001 tr.1–8; tài liệu 10. [CẦN ĐIỀN hồ sơ nghiệm thu và phê duyệt].'})
rowmap(d.tables[13],{
 'Định danh phân đoạn':'Đã ghi VLAN 10, 20, 30, 31, 40, 70, 100, 235 và NIC untagged; IP quản trị/iDRAC thuộc dải 10.3.99.x. [CẦN ĐIỀN VRF/VNI, network ID và xác nhận VLAN 99].',
 'Địa chỉ và điểm cuối':'IP thiết bị, iDRAC, IP VM từ control plane và 06 floating IP tại tài liệu 10. IP guest của VM Proxmox chưa được xác nhận.',
 'HA và phụ thuộc':'02 FortiGate trong HA; 02 Nexus; host có bond. [CẦN ĐIỀN peer, LB, DNS/NTP và phép thử failover].',
 'Hồ sơ hiện trạng':'EV-001 tr.3–5, 9–15; tài liệu 10. [CẦN ĐIỀN export route/ACL và ánh xạ vùng Z].'})
rowmap(d.tables[15],{
 'Nơi lưu và bản sao':'Ceph có pool .mgr, images, volumes, vms, backups; size 3/min_size 2; 19,20 TB raw. 26 volume, tổng 2.387 GiB logic. [CẦN ĐIỀN ánh xạ dữ liệu, quyền và bản sao ngoài cụm].'})
rowmap(d.tables[16],{
 'HA và quorum':'03 VM controller trên 03 Proxmox; Ceph 03 host/10 OSD, size 3/min_size 2; 02 FortiGate HA. [CẦN ĐIỀN quorum DB/queue, miền lỗi và kết quả thử].',
 'Sao lưu':'Pool backups ghi 0,00 GB logic tại thời điểm đọc; S3 ngoài site chưa kiểm kê. [CẦN ĐIỀN job, lịch, retention, mã hóa và kho độc lập].',
 'Năng lực DR và rủi ro chung':'mfa-02 ngoài site: 2 vCPU, RAM Linux 3,82 GiB, vda 80 GiB. [CẦN XÁC NHẬN IDC đặt VM; CẦN ĐIỀN năng lực DR toàn hệ thống và miền lỗi chung].',
 'Kết quả xác nhận':'EV-001 xác nhận số liệu kiểm kê; chưa có kết quả diễn tập. [CẦN ĐIỀN RPO/RTO đo được và biên bản phục hồi].'})
rowmap(d.tables[17],{
 'Thành phần và phạm vi':'09 host, 21 VM Proxmox, 24 VM compute tại site chính; 01 VM ngoài site; 02 switch và 02 firewall.',
 'Nguồn và bản dựng':'NX-OS 9.3(8); FortiOS 7.6.6 build3652. [CẦN ĐIỀN phiên bản XOR, Proxmox, OS, Ceph, commit/tag và digest].',
 'Bằng chứng và phê duyệt':'EV-001 tr.2–5, 9–15. [CẦN ĐIỀN manifest, SBOM, log triển khai và văn bản phê duyệt].'})
save(d,p)

# Reuse each requirement and plan verbatim; add only scoped inventory observations.
p=ADMIN/'02_Phu_luc_doi_chieu_yeu_cau_TCVN.docx';d=Document(p)
for t in d.tables:
 for r in t.rows[1:]:
  code=r.cells[0].text.split('\n')[0]
  if code in OBS:
   observation,gap=OBS[code]
   cell(r.cells[2],r.cells[2].text.replace('Hiện trạng: [CẦN ĐIỀN].','Hiện trạng: '+observation+' Còn thiếu: '+gap))
   cell(r.cells[3],r.cells[3].text.replace('Mã BC: [CẦN ĐIỀN].','Mã BC: EV-001 (kiểm kê 18/10/2026); tài liệu 10.'))
save(d,p)

p=TECH/'01_So_tay_ra_soat_ky_thuat.docx';d=Document(p)
replace(d,'Danh mục tài sản, dịch vụ thực tế, phiên bản, số lượng và ranh giới khách hàng.','EV-001 đã ghi 13 thiết bị chính, 46 VM; bổ sung vị trí, vòng đời, chủ tài sản, phiên bản và ranh giới khách hàng.')
replace(d,'chưa có hạng mục nào được xác nhận đạt chỉ từ việc lập bộ hồ sơ này.','đã cập nhật số liệu kiểm kê ngày 18/10/2026, chưa có kết luận đạt kiểm soát từ việc kiểm kê.')
replace(d,'Không có kết quả kiểm tra thực tế được đưa vào.','Có số liệu kiểm kê EV-001; kết quả kiểm thử an ninh mạng vẫn chờ xác nhận.')
save(d,p)

# Report remains an unsigned results form; inventory is distinguished from security tests.
p=ADMIN/'09_Bao_cao_ket_qua_ra_soat.docx';d=Document(p)
replace(d,'Thời gian: [CẦN ĐIỀN]. Phiên bản hệ thống và hồ sơ: [CẦN ĐIỀN].','Số liệu kiểm kê: 18/10/2026 theo EV-001; thời gian rà soát an ninh mạng: [CẦN ĐIỀN]. Phiên bản hệ thống: [CẦN ĐIỀN]; bộ hồ sơ 2.5.')
replace(d,'Thực hiện [CẦN ĐIỀN số, ngày kế hoạch hoặc văn bản giao nhiệm vụ], Khối Hạ tầng & Vận hành báo cáo kết quả rà soát như sau:','Thực hiện [CẦN ĐIỀN số, ngày kế hoạch hoặc văn bản giao nhiệm vụ], Khối Hạ tầng & Vận hành báo cáo kết quả rà soát như sau:')
save(d,p)

p=ADMIN/'00_Danh_muc_ho_so.docx';d=Document(p)
for t in d.tables:
 for r in t.rows:
  if any('10'==c.text.strip() for c in r.cells):
   for c in r.cells:
    if 'Phụ lục' in c.text:cell(c,'Phụ lục thiết kế và hiện trạng; kiểm kê ngày 18/10/2026 và sơ đồ tài sản')
save(d,p)

# Technical reconciliation incorporates the new evidence and keeps historical findings.
d=Doc('Đối soát hiện trạng kiểm kê và tồn tại kỹ thuật',False)
d.cover('Cập nhật kiểm kê ngày 18 tháng 10 năm 2026')
d.p('Khối Hạ tầng & Vận hành đối chiếu hồ sơ thiết kế v0.9.4 với '+SOURCE_TITLE+'. Số liệu kiểm kê được cập nhật vào checklist và phụ lục 10 của bộ hành chính; kết quả kiểm thử và phê duyệt được quản lý riêng.')
d.h('1 Nguồn và phạm vi');d.p(SCOPE);d.p(SUMMARY);d.p(LIMIT)
d.table(['Dữ liệu','Căn cứ và cách sử dụng'],[
 ('EV-001',SOURCE+'; 16 trang. SHA-256: '+SHA),
 ('H01–H06','Các tệp máy đọc được do báo cáo dẫn tại C:/XOR Secure/openstack-deployment/evidence; chưa được cung cấp trong lượt cập nhật này. Thu bản gốc và hash trước khi xác minh độc lập.'),
 ('129 mục lịch sử','Giữ nguyên dữ liệu tại Nguon_v094; ghi nhận kiểm kê vào cột bằng chứng/hành động của Doi_soat. Không chuyển kết quả lịch sử thành kết quả mới.'),
 ('119 mục TCVN','Giữ nguyên yêu cầu và kết luận chờ kiểm tra; chỉ thêm hiện trạng kiểm kê vào các mục liên quan.')],[1,3])
d.page('2 Sai khác giữa thiết kế và kiểm kê')
d.table(['Nội dung','Hồ sơ v0.9.4','Kiểm kê 18/10/2026'],[
 ('Controller','03 server vai trò Controller, 128 GB/node.','03 Proxmox vật lý: RAM 256/192/192 GiB; trên mỗi host có 01 VM controller 24 vCPU/48 GiB/300G local-lvm.'),
 ('Compute','03 server, 512 GB/node.','RAM DIMM 288/512/512 GiB; số VM 10/0/14. Không suy compute-02 lỗi hoặc là dự phòng đã kiểm chứng từ số VM bằng 0.'),
 ('Tổng RAM','2.304 GB theo quy hoạch.','2.336 GiB DIMM theo kiểm kê; 324 nhân/648 luồng CPU. Không cộng vCPU và RAM VM vào tổng vật lý.'),
 ('Ceph','03 server, 10 SSD 1,92 TB.','Đã ghi nhận 4/3/3 OSD, 19,20 TB raw; size 3/min_size 2; MAX AVAIL 5,02 TB theo pool tại thời điểm đọc, không cộng giữa các pool.'),
 ('SAN và switch SAN','SAN 24 × 600 GB; 14,4 TB raw; 01 switch SAN.','Chưa có model/serial/dung lượng xác minh; chưa thấy FC trên 09 host trong dữ liệu đã thu. Không đưa SAN vào tổng đã kiểm kê.'),
 ('DR và S3','Phương án phục hồi độc lập.','Mới kiểm kê mfa-02 ngoài site; chưa đủ dữ liệu cụm DR, S3 và đồng bộ.'),
 ('VLAN','10/20/30/31/40/70/99/100 theo thiết kế.','Đã có cấu hình VLAN/NIC VM; bổ sung VLAN 235 trunk 40 và NIC untagged. IP 10.3.99.x không tự xác nhận VLAN 99.')],[.8,1.45,2.7])
d.page('3 Công việc cần hoàn thiện từ kiểm kê')
d.table(['Mã và đầu mối','Ghi nhận','Hành động'],[(g[0]+'\n'+g[3],g[2],g[4]) for g in GAPS[:4]],[1,2.2,2.4])
d.page('4 Công việc lưu trữ phục hồi và an ninh')
d.table(['Mã và đầu mối','Ghi nhận','Hành động'],[(g[0]+'\n'+g[3],g[2],g[4]) for g in GAPS[4:]],[1,2.2,2.4])
d.p('Người phụ trách, hạn hoàn thành và kết quả: nhập tại checklist. Không điền ngày hoặc người xác minh khi mới nhận báo cáo; đây chưa phải biên bản nghiệm thu.')
d.page('5 Bàn giao và truy vết')
d.p('Danh mục Tai san có 59 bản ghi: 13 thiết bị chính và 46 VM. Mã KK là mã hồ sơ kiểm kê; UUID VM nền tảng, chủ sở hữu và thuộc tính còn thiếu tiếp tục được đối soát. Bộ điều khiển iDRAC và 03 controller VM không được cộng thành server vật lý.')
d.p('21 VM Proxmox có 166 vCPU và 340 GiB RAM cấu hình; 24 VM compute có 113 vCPU và 215,5 GiB RAM. Đây là tài nguyên cấp phát, không phải tải tiêu thụ. Mức RAM 3,82 GiB của mfa-02 là Linux đọc được, không cộng như RAM DIMM.')
d.p('Ba sơ đồ được lưu kèm: sơ đồ 01–02 là phương án mục tiêu; sơ đồ 03 là cấu trúc tài sản theo kiểm kê. Sơ đồ 03 không thay thế bản vẽ cáp hoặc chứng minh đồng bộ DC–DR.')
d.p('Số liệu được chuyển sang thuyết minh, phụ lục 10 và các trường hiện trạng có liên quan. Chuyên trách an ninh mạng xác nhận kết quả sau khi thu cấu hình và thực hiện phép thử theo phạm vi đã duyệt. Các phát hiện cũ về khóa trong image, registry và bản vá vẫn cần kiểm tra riêng; báo cáo kiểm kê không đóng các mục đó.')
d.save(TECH/'03_Ke_hoach_doi_soat_nguon_v094.docx');changed.append(str(TECH/'03_Ke_hoach_doi_soat_nguon_v094.docx'))

exec((TMP/'build_annex.py').read_text(),globals())
changed.append(str(ADMIN/'10_Phu_luc_thiet_ke_va_hien_trang.docx'))
(TMP/'changed-docs.json').write_text(json.dumps(changed,ensure_ascii=False,indent=2))
old_manifest=json.loads((ROOT/'tmp/ho-so-v2-4/doc_manifest.json').read_text())
(TMP/'doc_manifest.json').write_text(json.dumps([s.replace('CD4-v2.4','CD4-v2.5') for s in old_manifest],ensure_ascii=False,indent=2))
print('Updated',len(changed),'Word files')
