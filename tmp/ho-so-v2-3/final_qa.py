from pathlib import Path
from collections import Counter
import re,json,hashlib,zipfile
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

root=Path('/Users/thangchu/Documents/ChatGPT/XOR Secure Cloud')
tmp=Path(__file__).parent
old=root/'output/XOR-Ho-so-CD4-v2.2';new=root/'output/XOR-Ho-so-CD4-v2.3'
def norm(s):return re.sub(r'\s+',' ',s.replace('2.2','2.3').replace('Ô vàng','Ô nền xám nhạt').replace('trường vàng','trường nền xám nhạt')).strip().casefold()
docs=[]
for f in new.rglob('*.docx'):
 a=Document(old/f.relative_to(new));b=Document(f)
 ta=[norm(c.text) for t in a.tables for row in t.rows for c in row.cells]
 tb=[norm(c.text) for t in b.tables for row in t.rows for c in row.cells]
 assert ta==tb, f'Table content changed: {f.name}'
 pa=Counter(norm(p.text) for p in a.paragraphs if p.text.strip())
 pb=Counter(norm(p.text) for p in b.paragraphs if p.text.strip())
 diff=list((pa-pb).elements())
 allowed=('bản dự thảo 2.3 ngày','dc tại trung tâm hạ tầng')
 unapproved=[x for x in diff if not x.startswith(allowed)]
 assert not unapproved,(f.name,unapproved)
 assert b.styles['Normal'].font.name=='Times New Roman'
 docs.append({'file':f.name,'tableCellsPreserved':len(ta),'coverDifferences':diff})
books=[]
for f in new.rglob('*.xlsx'):
 a=load_workbook(old/f.relative_to(new));b=load_workbook(f);cache=load_workbook(f,data_only=True)
 assert a.sheetnames==b.sheetnames
 changes=[];formulas=0;errors=[]
 for sa in a:
  sb=b[sa.title]
  for row in sa:
   for ca in row:
    cb=sb[ca.coordinate]
    if ca.value!=cb.value:
     changes.append([sa.title,ca.coordinate,ca.value,cb.value])
     assert (isinstance(ca.value,str) and ca.value.replace('Ô vàng','Ô nền xám nhạt')==cb.value) or (sa.title=='Huong dan' and ca.coordinate=='B23'),changes[-1]
    if ca.data_type=='f':
     formulas+=1;assert ca.value==cb.value
    if cache[sa.title][ca.coordinate].data_type=='e':errors.append((sa.title,ca.coordinate))
  av=[(str(v.sqref),v.type,v.formula1,v.formula2) for v in sa.data_validations.dataValidation]
  bv=[(str(v.sqref),v.type,v.formula1,v.formula2) for v in sb.data_validations.dataValidation]
  assert av==bv,(sa.title,'validations')
 assert not errors,errors
 books.append({'file':f.name,'sheets':len(b.sheetnames),'formulasPreserved':formulas,'changes':changes,'formulaErrors':errors})
sources=[]
for p,expected in [('/Users/thangchu/Downloads/De an_xCommandCenter_v2 (1).pdf','5355b56aa0ff210052747f6456ba71f234daef5708533281832e15794c839e2b'),('/Users/thangchu/Work/XOR/Xor Secure Cloud/XOR_SECURE_CLOUD_Ho_so_Cap_do_4_v0.9.4.docx','ec28ac067dba4d046d7586798dc38c744e7c09c71ca25dca48883f334ebdc5c5')]:
 sha=hashlib.sha256(Path(p).read_bytes()).hexdigest();assert sha==expected
 sources.append({'file':p,'sha256':sha,'unchanged':True})
pages=[]
for x in json.loads((tmp/'render-results.json').read_text()):
 pdf=next(Path(x['dir']).glob('*.pdf'));reader=PdfReader(pdf)
 assert len(reader.pages)==x['pages'],x
 for i,p in enumerate(reader.pages):assert len(p.extract_text().strip())>5,(x['file'],i+1,'blank page')
 pages.append({'file':x['file'],'pages':len(reader.pages)})
qa={'docx':docs,'xlsx':books,'sources':sources,'pages':pages,'totalPages':sum(x['pages'] for x in pages),'visualReview':'All first-pass page layouts inspected in 19 contact sheets; dense pages and modified covers/sections inspected at full size. All 12 workbook tabs and SVG/PNG diagram inspected.'}
(tmp/'final-qa.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))
print(json.dumps({'documents':len(docs),'workbooks':len(books),'pages':qa['totalPages'],'valueChanges':sum(len(x['changes']) for x in books),'formulaCounts':[x['formulasPreserved'] for x in books]},ensure_ascii=False))
