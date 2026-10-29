from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def restyle_doc_class(Base):
 class StyledDoc(Base):
  def __init__(self,title,formal=True):
   super().__init__(title,formal)
   self.font='Times New Roman';self.size=13
   for sec in self.doc.sections:
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.left_margin=Cm(3);sec.right_margin=Cm(2);sec.top_margin=sec.bottom_margin=Cm(2)
    sec.header_distance=Cm(.7);sec.footer_distance=Cm(1)
   self.width=self.doc.sections[0].page_width-self.doc.sections[0].left_margin-self.doc.sections[0].right_margin
   for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    st=self.doc.styles[name];st.font.name=self.font;st.font.size=Pt(13);st.font.color.rgb=RGBColor(0,0,0)
    f=st.element.get_or_add_rPr().get_or_add_rFonts()
    for a in ['ascii','hAnsi','eastAsia','cs']:f.set(qn('w:'+a),self.font)
    st.paragraph_format.first_line_indent=Cm(1)
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    st.paragraph_format.line_spacing=1.5
    st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(3)
   for name in ['Heading 1','Heading 2','Heading 3']:
    st=self.doc.styles[name];st.font.bold=True;st.paragraph_format.first_line_indent=Cm(0)
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.keep_with_next=True
   for name in ['Title','Subtitle','Caption']:
    st=self.doc.styles[name];st.paragraph_format.first_line_indent=Cm(0);st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
   self.doc.styles['Title'].font.size=Pt(15);self.doc.styles['Title'].font.bold=True
   self.doc.styles['Subtitle'].font.italic=True
   self.doc.styles['Caption'].font.size=Pt(12);self.doc.styles['Caption'].font.italic=True
   self.doc.styles['Caption'].paragraph_format.line_spacing=1.15
   h=self.doc.sections[0].header.paragraphs[0];h.text='';h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
   h.paragraph_format.first_line_indent=Cm(0)
   for r in h.runs:r.font.name=self.font;r.font.size=Pt(8);r.italic=True;r.font.color.rgb=RGBColor(0,0,0)
   self._footer(self.doc.sections[0])
  def _footer(self,sec):
   p=sec.footer.paragraphs[0];p.clear();p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
   r=p.add_run();r.font.name=self.font;r.font.size=Pt(10)
   f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.addnext(f)
  def p(self,text='',bold=False,style=None,align=None):
   p=super().p(text,bold,style,align)
   if align is not None:p.paragraph_format.first_line_indent=Cm(0)
   if not text:p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0)
   return p
  def h(self,text,level=2):
   return super().h(text.upper() if level==1 else text,level)
  def cover(self,subtitle):
   sec=self.doc.sections[0];sec.header.paragraphs[0].clear();sec.footer.paragraphs[0].clear()
   def c(text,size=13,bold=False,italic=False,before=0,after=0,style=None):
    p=self.doc.add_paragraph(style=style);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.line_spacing=1.15;pf.space_before=Pt(before);pf.space_after=Pt(after)
    r=p.add_run(text);r.font.name=self.font;r.font.size=Pt(size);r.bold=bold;r.italic=italic;r.font.color.rgb=RGBColor(0,0,0)
    return p
   c('CÔNG TY CỔ PHẦN XOR CLOUD',13,True)
   p=c('KHỐI HẠ TẦNG & VẬN HÀNH',13,True,after=18)
   line=c('',after=72);line.paragraph_format.left_indent=Cm(4.4);line.paragraph_format.right_indent=Cm(4.4)
   borders=OxmlElement('w:pBdr');bottom=OxmlElement('w:bottom');bottom.set(qn('w:val'),'single');bottom.set(qn('w:sz'),'6');bottom.set(qn('w:space'),'1');bottom.set(qn('w:color'),'000000');borders.append(bottom);line._p.get_or_add_pPr().append(borders)
   category='HỒ SƠ' if self.formal else 'TÀI LIỆU KỸ THUẬT'
   c(category,22,True,after=13)
   titles={
    'Sổ tay rà soát kỹ thuật hệ thống thông tin cấp độ 4':'SỔ TAY RÀ SOÁT KỸ THUẬT\nHỆ THỐNG THÔNG TIN CẤP ĐỘ 4',
    'Hướng dẫn hoàn thiện bộ hồ sơ đề xuất cấp độ 4':'HƯỚNG DẪN HOÀN THIỆN\nBỘ HỒ SƠ ĐỀ XUẤT CẤP ĐỘ 4',
    'Hồ sơ đề xuất cấp độ 4 hệ thống XOR SECURE CLOUD':'HỒ SƠ ĐỀ XUẤT CẤP ĐỘ 4\nHỆ THỐNG XOR SECURE CLOUD',
    'Phụ lục đối chiếu yêu cầu và phương án bảo đảm an ninh mạng':'PHỤ LỤC ĐỐI CHIẾU YÊU CẦU\nVÀ PHƯƠNG ÁN BẢO ĐẢM AN NINH MẠNG',
    'Quy chế và quy trình bảo đảm an ninh mạng XOR SECURE CLOUD':'QUY CHẾ VÀ QUY TRÌNH\nBẢO ĐẢM AN NINH MẠNG XOR SECURE CLOUD',
    'Mẫu ý kiến chuyên môn và báo cáo tự rà soát':'MẪU Ý KIẾN CHUYÊN MÔN\nVÀ BÁO CÁO TỰ RÀ SOÁT',
    'Mẫu thông báo cấp độ và đề nghị đánh giá độc lập':'MẪU THÔNG BÁO CẤP ĐỘ\nVÀ ĐỀ NGHỊ ĐÁNH GIÁ ĐỘC LẬP',
    'Kế hoạch đối soát thiết kế và tồn tại từ hồ sơ tham khảo':'KẾ HOẠCH ĐỐI SOÁT THIẾT KẾ\nVÀ TỒN TẠI TỪ HỒ SƠ THAM KHẢO',
    'Phụ lục đối chiếu thiết kế và xác nhận hiện trạng hệ thống':'PHỤ LỤC ĐỐI CHIẾU THIẾT KẾ\nVÀ XÁC NHẬN HIỆN TRẠNG HỆ THỐNG',
   }
   c(titles.get(self.title,self.title.upper()),15,True,after=14,style='Title')
   c('XOR SECURE CLOUD',18,True,after=12)
   c(subtitle,13,italic=True,after=12)
   c('Nền tảng XOR STACK CLOUD',13,after=56)
   c('Đơn vị vận hành: Khối Hạ tầng & Vận hành',13,after=6)
   c('DC: Trung tâm Hạ tầng IDC Cầu Giấy\nDR: Trung tâm Hạ tầng IDC Nam Thăng Long',13,after=10)
   c('',12,after=42)
   c('Hà Nội, tháng 10 năm 2026',13,italic=True)
   body=self.doc.add_section(WD_SECTION_START.NEW_PAGE)
   body.header.is_linked_to_previous=False;body.footer.is_linked_to_previous=False
   body.header.paragraphs[0].clear();self._footer(body)
   n=OxmlElement('w:pgNumType');n.set(qn('w:start'),'1');body._sectPr.append(n)
  def table(self,headers,rows,widths=None,size=None):
   t=super().table(headers,rows,widths,10.5)
   for i,row in enumerate(t.rows):
    for c in row.cells:
     tc=c._tc.get_or_add_tcPr()
     for el in list(tc.findall(qn('w:shd'))):tc.remove(el)
     sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'FFFFFF');tc.append(sh)
     for bd in tc.findall(qn('w:tcBorders')):
      for edge in bd:edge.set(qn('w:color'),'000000');edge.set(qn('w:sz'),'4')
     for p in c.paragraphs:
      pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.line_spacing=1;pf.space_before=Pt(1);pf.space_after=Pt(2)
      p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
      for r in p.runs:r.font.name=self.font;r.font.size=Pt(10.5);r.font.color.rgb=RGBColor(0,0,0);r.bold=i==0
   return t
  def sign(self,*a,**kw):
   t=super().sign(*a,**kw)
   for i,c in enumerate(t.rows[0].cells):
    for p in c.paragraphs:
     p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.15
     for r in p.runs:r.font.name=self.font;r.font.size=Pt(12 if i==0 else 13);r.bold=i==1
   return t
  def letterhead(self,sender,ref,subject=None,model=None):
   super().letterhead(sender.upper(),ref,subject,None)
   t=self.doc.tables[-1]
   for i,c in enumerate(t.rows[0].cells):
    value=c.text;c.text=''
    lines=value.split('\n')
    for j,line in enumerate(lines):
     p=c.paragraphs[0] if j==0 else c.add_paragraph()
     p.alignment=WD_ALIGN_PARAGRAPH.CENTER
     p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(3)
     r=p.add_run(line);r.font.name=self.font;r.font.size=Pt(12)
     r.bold=j==0 or (i==1 and j==1)
     if line.startswith('Hà Nội,'):r.italic=True;r.font.size=Pt(13)
  def topology(self,prefix='2'):
   for i,name in enumerate(['01_Topo_tong_the_DC_DR','02_Topo_phan_vung_va_luong']):
    sec=self.doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec.page_width=Cm(29.7);sec.page_height=Cm(21)
    sec.left_margin=sec.right_margin=Cm(2);sec.top_margin=sec.bottom_margin=Cm(1.5)
    sec.orientation=1
    for n in list(sec._sectPr.findall(qn('w:pgNumType'))):sec._sectPr.remove(n)
    self.h(prefix+(' ' if prefix=='Phụ lục' else '.')+str(i+1)+' '+('Phương án kiến trúc tổng thể DC–DR' if i==0 else 'Phương án phân vùng mạng và kết nối'),2)
    p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0)
    p.add_run().add_picture(str(__import__('pathlib').Path(__file__).parent/(name+'.png')),width=Cm(25.3))
    cap=self.p('Hình '+str(i+1)+'. '+('Kiến trúc tổng thể hệ thống XOR SECURE CLOUD.' if i==0 else 'Phân vùng mạng và các luồng kết nối chính.'),style='Caption')
    cap.paragraph_format.space_after=Pt(0)
   if not self.formal:return
   sec=self.doc.add_section(WD_SECTION_START.NEW_PAGE)
   sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.orientation=0
   sec.left_margin=Cm(3);sec.right_margin=Cm(2);sec.top_margin=sec.bottom_margin=Cm(2)
   for n in list(sec._sectPr.findall(qn('w:pgNumType'))):sec._sectPr.remove(n)
  def save(self,path):
   # Keep the dense platform subsection together without reducing font size.
   if self.title=='Hồ sơ đề xuất cấp độ 4 hệ thống XOR SECURE CLOUD':
    compact=False
    for p in self.doc.paragraphs:
     if p.text.startswith('4 PHƯƠNG ÁN NỀN TẢNG'):compact=True
     if p.text.startswith('PHẦN III'):compact=False
     if compact:
      p.paragraph_format.line_spacing=1.35
      p.paragraph_format.space_before=Pt(6 if p.style.name.startswith('Heading') else 0)
      p.paragraph_format.space_after=Pt(3)
   # The ninth control in this table must not occupy a separate page alone.
   for t in self.doc.tables:
    if len(t.rows)>1 and t.cell(1,0).text.startswith('TC-089'):
     spacer=t._tbl.getnext()
     if spacer is not None and spacer.tag==qn('w:p') and not ''.join(spacer.itertext()).strip():
      spacer.getparent().remove(spacer)
     for row in t.rows:
      for cell in row.cells:
       for side in ['top','bottom']:
        for node in cell._tc.iter(qn('w:'+side)):
         if node.getparent().tag==qn('w:tcMar'):node.set(qn('w:w'),'40')
       for p in cell.paragraphs:
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
   return super().save(path)
 return StyledDoc
