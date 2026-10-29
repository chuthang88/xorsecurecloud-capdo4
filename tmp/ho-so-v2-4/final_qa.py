from pathlib import Path
from docx import Document
from zipfile import ZipFile
from collections import Counter
import hashlib,re,json,openpyxl
from copy import copy
P=Path(__file__).parent;root=P.parents[1];old=root/'output/XOR-Ho-so-CD4-v2.3';new=root/'output/XOR-Ho-so-CD4-v2.4'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
qa={'source_files':{},'docx':[],'xlsx':{}}
for src,expected in [('/Users/thangchu/Work/XOR/Xor Secure Cloud/XOR_SECURE_CLOUD_Ho_so_Cap_do_4_v0.9.4.docx','ec28ac067dba4d046d7586798dc38c744e7c09c71ca25dca48883f334ebdc5c5'),('/Users/thangchu/Downloads/De an_xCommandCenter_v2 (1).pdf','5355b56aa0ff210052747f6456ba71f234daef5708533281832e15794c839e2b')]:
 v=sha(Path(src));assert v==expected;qa['source_files'][src]=v
for path in new.rglob('*.docx'):
 d=Document(path);text='\n'.join([x.text for x in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
 for sec in d.sections:text+='\n'+'\n'.join(p.text for p in sec.header.paragraphs)
 assert 'dựa trên OpenStack' not in text
 assert 'OpenStack' not in text,(path.name,'OpenStack')
 if 'Bo-2' in str(path):
  for term in ['Dự thảo','dự thảo','Mẫu nội bộ','MẪU','Hướng dẫn hoàn thiện','workbook','chưa ký ban hành','KẾT LUẬN THẬT']:
   assert term not in text,(path.name,term)
 qa['docx'].append({'file':str(path.relative_to(new)),'paragraphs':len(d.paragraphs),'tables':len(d.tables),'images':len(d.inline_shapes)})
a=Document(old/'Bo-2-Hanh-chinh/02_Phu_luc_doi_chieu_yeu_cau_TCVN.docx');b=Document(new/'Bo-2-Hanh-chinh/02_Phu_luc_doi_chieu_yeu_cau_TCVN.docx')
def controls(d):
 return {row.cells[0].text:(row.cells[1].text,row.cells[2].text,row.cells[3].text) for t in d.tables for row in t.rows if re.match('TC-\\d{3}\\n',row.cells[0].text)}
x,y=controls(a),controls(b);assert x.keys()==y.keys() and len(y)==119
assert all(x[k][0]==y[k][0] for k in x)
qa['controls']={'retained':len(y),'requirements_identical':True,'wording_changes':[k for k in x if x[k]!=y[k]]}
fn='04_Doi_soat_129_hang_muc_v094.xlsx';assert sha(old/'Bo-1-Ky-thuat'/fn)==sha(new/'Bo-1-Ky-thuat'/fn)
qa['source_129_workbook_unchanged']=True
fn='02_Checklist_ra_soat_XOR_cap_do_4.xlsx';a=openpyxl.load_workbook(old/'Bo-1-Ky-thuat'/fn);b=openpyxl.load_workbook(new/'Bo-1-Ky-thuat'/fn)
assert a.sheetnames==b.sheetnames
changes=[];formula_count=0;validation_count=0
for sa,sb in zip(a,b):
 assert (sa.max_row,sa.max_column)==(sb.max_row,sb.max_column),(sa.title,'extent')
 assert sa.freeze_panes==sb.freeze_panes,(sa.title,'panes')
 assert set(sa.tables)==set(sb.tables),(sa.title,'tables')
 assert [str(x) for x in sa.merged_cells]==[str(x) for x in sb.merged_cells]
 da=[(str(d.sqref),d.type,d.formula1) for d in sa.data_validations.dataValidation];db=[(str(d.sqref),d.type,d.formula1) for d in sb.data_validations.dataValidation]
 assert da==db,(sa.title,'validation');validation_count+=len(db)
 for ra,rb in zip(sa,sb):
  for ca,cb in zip(ra,rb):
   if ca.data_type=='f':assert ca.value==cb.value;formula_count+=1
   if ca.value!=cb.value:changes.append({'sheet':sa.title,'cell':ca.coordinate,'before':ca.value,'after':cb.value})
   assert copy(ca.font)==copy(cb.font),(sa.title,ca.coordinate,'font')
   assert copy(ca.fill)==copy(cb.fill),(sa.title,ca.coordinate,'fill')
   assert ca.number_format==cb.number_format
allowed={('Cloud','D6'),('Huong dan','E31'),('Ho so','E5'),('Ho so','E10'),('Ho so','E14'),('Ho so','E20'),('Ho so','E22'),('Ho so','E23')}|{('Cloud',f'B{i}') for i in range(6,24)}
assert all((c['sheet'],c['cell']) in allowed for c in changes)
ac=openpyxl.load_workbook(old/'Bo-1-Ky-thuat'/fn,data_only=True);bc=openpyxl.load_workbook(new/'Bo-1-Ky-thuat'/fn,data_only=True)
for sa,sb in zip(ac,bc):
 for ra,rb in zip(sa,sb):
  for ca,cb in zip(ra,rb):
   if (sa.title,ca.coordinate) not in allowed:assert ca.value==cb.value,(sa.title,ca.coordinate,'cached')
qa['xlsx']={'changed_cells':changes,'unchanged_formulas':formula_count,'preserved_validations':validation_count,'cached_results_unchanged':True}
# All revised assets exist in each packet.
for folder in ['Bo-1-Ky-thuat','Bo-2-Hanh-chinh']:
 for name in ['01_Topo_tong_the_DC_DR','02_Topo_phan_vung_va_luong']:
  for ext in ['svg','png']:assert (new/folder/(name+'.'+ext)).stat().st_size>1000
qa['render_results']=json.loads((P/'render-results.json').read_text())
assert all(r['code']==0 and r['pages']>0 for r in qa['render_results'])
(P/'final-qa.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))
print(json.dumps({'documents':len(qa['docx']),'controls':qa['controls'],'pages':sum(r['pages'] for r in qa['render_results']),'excel_changed_cells':len(changes),'formulas_retained':formula_count,'validation_rules_retained':validation_count,'source_files_unchanged':True},ensure_ascii=False))
