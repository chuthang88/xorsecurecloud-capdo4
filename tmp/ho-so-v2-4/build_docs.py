from pathlib import Path
from copy import deepcopy
import json, shutil
from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from content import *

TECH=OUT/'Bo-1-Ky-thuat'; ADMIN=OUT/'Bo-2-Hanh-chinh'
TECH.mkdir(parents=True,exist_ok=True);ADMIN.mkdir(parents=True,exist_ok=True)
CREATED=[]
DRAFT=''

class Doc:
 def __init__(self,title,formal=True):
  self.doc=Document();self.formal=formal;self.title=title
  self.font='Times New Roman' if formal else 'Arial';self.size=13 if formal else 11
  s=self.doc.sections[0]
  if formal:s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=Cm(3);s.right_margin=Cm(2);s.top_margin=s.bottom_margin=Cm(2)
  else:s.page_width=Inches(8.5);s.page_height=Inches(11);s.left_margin=s.right_margin=Inches(.7);s.top_margin=s.bottom_margin=Inches(.65)
  self.width=s.page_width-s.left_margin-s.right_margin
  s.header_distance=s.footer_distance=Cm(.7)
  for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
   st=self.doc.styles[name];st.font.name=self.font;st.font.size=Pt(self.size);st.font.color.rgb=RGBColor(0,0,0)
   fonts=st.element.get_or_add_rPr().get_or_add_rFonts()
   for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
    if qn('w:'+attr) in fonts.attrib:del fonts.attrib[qn('w:'+attr)]
   for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),self.font)
   st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.1
   for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
  for name,size in [('Title',19 if formal else 23),('Heading 1',15),('Heading 2',13),('Heading 3',12)]:
   st=self.doc.styles[name];st.font.size=Pt(size);st.font.bold=True
  self.doc.styles['Caption'].font.size=Pt(10)
  header=s.header.paragraphs[0];header.text=DRAFT
  for r in header.runs:r.font.name=self.font;r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
  f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
  f.add_run('XOR SECURE CLOUD  •  Phiên bản 2.4  •  ').font.size=Pt(9)
  fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
  self.doc.core_properties.title=title;self.doc.core_properties.author='XOR CLOUD';self.doc.core_properties.subject='Hồ sơ lần đầu đề xuất cấp độ 4'
 def p(self,text='',bold=False,style=None,align=None):
  p=self.doc.add_paragraph(style=style);p.add_run(text).bold=True if style=='Title' else bold
  if align is not None:p.alignment=align
  return p
 def h(self,text,level=2):return self.doc.add_heading(text,level)
 def page(self,title):
  p=self.h(title,1);p.paragraph_format.page_break_before=True;return p
 def cover(self,subtitle):
  self.p('XOR SECURE CLOUD',True,align=WD_ALIGN_PARAGRAPH.CENTER)
  self.p(self.title,style='Title',align=WD_ALIGN_PARAGRAPH.CENTER)
  self.p(subtitle,align=WD_ALIGN_PARAGRAPH.CENTER)
  self.p('Nền tảng XOR STACK CLOUD',align=WD_ALIGN_PARAGRAPH.CENTER)
  self.p('DC tại Trung tâm Hạ tầng IDC Cầu Giấy\nDR tại Trung tâm Hạ tầng IDC Nam Thăng Long',align=WD_ALIGN_PARAGRAPH.CENTER)
  self.p('Phiên bản 2.4 ngày 17 tháng 10 năm 2026',align=WD_ALIGN_PARAGRAPH.CENTER)
 def table(self,headers,rows,widths=None,size=None):
  size=size or (11.5 if self.formal else 10)
  t=self.doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
  if widths is None:widths=[1]*len(headers)
  scale=self.width/sum(widths)
  for i,w in enumerate(widths):t.columns[i].width=int(w*scale)
  for i,x in enumerate(headers):t.rows[0].cells[i].text=str(x)
  for row in rows:
   for c,x in zip(t.add_row().cells,row):c.text=str(x)
  for ri,row in enumerate(t.rows):
   tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
   if ri==0:tr.append(OxmlElement('w:tblHeader'))
   for ci,c in enumerate(row.cells):
    c.width=int(widths[ci]*scale);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9ECEF' if self.formal and ri==0 else ('203E59' if ri==0 else ('F5F7F9' if ri%2 else 'FFFFFF')));tc.append(sh)
    mar=OxmlElement('w:tcMar')
    for side in ['top','left','bottom','right']:
     e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
    tc.append(mar);bs=OxmlElement('w:tcBorders')
    for side in ['top','bottom','left','right']:
     b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'BFC6CC');bs.append(b)
    tc.append(bs)
    for p in c.paragraphs:
     p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.03
     for r in p.runs:
      r.font.name=self.font;r.font.size=Pt(size);r.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 and not self.formal else '000000')
  self.p('').paragraph_format.space_after=Pt(1)
  return t
 def sign(self,recipient='Như trên;\nLưu: hồ sơ XOR SECURE CLOUD.',role=ROLE,signer=SIGNER):
  for p in self.doc.paragraphs[-2:]:p.paragraph_format.keep_with_next=True
  t=self.doc.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=int(self.width*.4);t.columns[1].width=int(self.width*.6)
  t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  t.cell(0,0).text='Nơi nhận:\n'+recipient;t.cell(0,1).text=role+'\n\n\n\n'+signer
  for i,c in enumerate(t.rows[0].cells):
   for p in c.paragraphs:
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.name=self.font;r.font.size=Pt(11)
  return t
 def letterhead(self,sender,ref,subject=None,model=None):
  if model:self.p(model,align=WD_ALIGN_PARAGRAPH.RIGHT)
  t=self.doc.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=int(self.width*.4);t.columns[1].width=int(self.width*.6)
  t.cell(0,0).text=sender+'\nSố: [CẦN ĐIỀN]/'+ref
  t.cell(0,1).text='CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n\nHà Nội, ngày [  ] tháng [  ] năm [  ]'
  for c in t.rows[0].cells:
   for p in c.paragraphs:
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(5)
    for r in p.runs:r.font.name=self.font;r.font.size=Pt(11.5);r.bold=True
  if subject:self.p(subject,True)
 def save(self,path):self.doc.save(path);CREATED.append(str(path));return path

def sources(d,cloud=False):
 d.h('Tài liệu tham chiếu')
 d.p('[S1] Nghị định 331/2026/NĐ-CP: Điều 14, 18, 20–24 và Phụ lục. Nguồn chính thức: '+URL)
 d.p('[S2] TCVN 14423:2026: Điều 6, trang 58–81; Phụ lục A.4, trang 109–111.')
 if cloud:d.p('Dự thảo TCVN Cloud năm 2026 chỉ dùng để rà soát thiết kế bổ sung. Các điểm bổ sung không tạo kết luận Hạng B/Mức 4 và không thay checklist toàn bộ dự thảo.')

from doc_style import restyle_doc_class
Doc=restyle_doc_class(Doc)

# BỘ 1
d=Doc('Sổ tay rà soát kỹ thuật hệ thống thông tin cấp độ 4',False)
d.cover('Bộ 1 dành cho đội kỹ thuật và các đầu mối cung cấp bằng chứng')
d.p(ORG+'. Đơn vị vận hành: '+OPERATOR+'. Đầu mối lập hồ sơ: '+AUTHOR_NAME+' — '+AUTHOR_ROLE+'.')
d.p('Mục tiêu là xác định rõ kiểm soát nào đã triển khai, bằng chứng nào đã có, nội dung nào chưa đủ và ai chịu trách nhiệm bổ sung. Dùng cùng workbook 02_Checklist_ra_soat_XOR_cap_do_4.xlsx; chưa có hạng mục nào được xác nhận đạt chỉ từ việc lập bộ hồ sơ này.')
d.table(['Đầu vào đã xác nhận','Dữ liệu phải bổ sung'],[
 ('Toàn bộ nền tảng cloud và dịch vụ quản trị chung.','Danh mục tài sản, dịch vụ thực tế, phiên bản, số lượng và ranh giới khách hàng.'),
 ('Nền tảng XOR STACK CLOUD.','Upstream release, bản dựng XOR, SBOM, cấu hình và kết quả thử.'),
 ('DC Cầu Giấy; DR Nam Thăng Long; lập hồ sơ lần đầu.','Hợp đồng và tài liệu IDC; phê duyệt thiết kế; mã số doanh nghiệp, thẩm quyền ký và liên hệ.')],[1,1])
d.h('Sản phẩm của một đợt rà soát')
d.p('Workbook đã cập nhật; kho bằng chứng có mã, phiên bản và người xác minh; danh sách thiếu sót có đầu mối và hạn; biên bản kiểm thử; báo cáo kết quả có giới hạn phạm vi. Chỉ chuyển số liệu đã xác nhận sang bộ hồ sơ hành chính.')
d.page('1 Cách sử dụng checklist')
d.table(['Bước','Thực hiện','Điều kiện kết thúc'],[
 ('1 Chốt phạm vi','Điền phạm vi từng dòng, dịch vụ và hai IDC. Chốt phần Cloud tham khảo tự nguyện.','Không dùng “không áp dụng” để che hạng mục chưa có.'),
 ('2 Giao việc','Gán cá nhân thực hiện từ đầu mối đề xuất; nhập hạn xử lý.','Mỗi mục có người chịu trách nhiệm rõ ràng.'),
 ('3 Kiểm tra hiện trạng','Đối chiếu cấu hình, tài liệu, log, quan sát và phép thử.','Phân biệt Chưa xác minh, Chưa triển khai, Một phần, Đầy đủ.'),
 ('4 Ghi bằng chứng','Đăng ký mã EV, kho lưu, ngày, phiên bản, phạm vi và người thu.','Bằng chứng đọc được, khớp môi trường và không lộ bí mật.'),
 ('5 Xác nhận','Người kiểm tra độc lập với người thực hiện khi phù hợp ghi kết quả/ngày.','Đạt chỉ khi đáp ứng đủ các ý trong dòng và phạm vi.'),
 ('6 Khắc phục','Lập hành động, thay đổi theo quy trình và retest.','Giữ kết quả cũ và mới; không xóa dấu vết chưa đạt.')],[.8,2.1,1.7])
d.h('Ý nghĩa trạng thái tổng hợp')
d.p('“Đủ hồ sơ kiểm tra” cần đồng thời: áp dụng Có, triển khai Đầy đủ, có mã hoặc đường dẫn bằng chứng, kết quả Đạt, có người và ngày kiểm tra. Công thức chỉ kiểm tra dữ liệu khai báo; người xác nhận chịu trách nhiệm về nội dung bằng chứng. Trạng thái này không phải quyết định phê duyệt cấp độ.')
d.p('“Không áp dụng đã duyệt” cần lý do và văn bản/người có thẩm quyền chấp thuận, người xác nhận và ngày. Nếu chưa chốt phạm vi hoặc chưa có bằng chứng, giữ nguyên trạng thái chờ, không ghi đạt.')
d.page('2 Phân công và ranh giới rà soát')
d.table(['Nhóm','Chịu trách nhiệm','Bằng chứng chính'],[
 ('Chủ quản và GRC','Phạm vi, căn cứ cấp độ, rủi ro, thẩm quyền và hồ sơ trình.','HS01–HS10, BIA, sổ rủi ro.'),
 ('Platform và DevSecOps','Bản dựng, control plane, hypervisor, image và pipeline.','CMDB, SBOM, baseline, release, T02/T07/T10/T11.'),
 ('Network và Storage','Phân vùng, kết nối, dữ liệu, HA và hiệu năng.','LLD, rule, route, ACL, replica, T01/T02/T04/T07.'),
 ('IAM và quản trị khóa','Định danh, MFA, đặc quyền, khóa và chứng thư.','RBAC, PAM, vòng đời khóa, T03/T04.'),
 ('SOC và CSIRT','Log, trực giám sát, sự cố, EDR/DLP và diễn tập.','Ma trận nguồn, lịch trực, ticket, T05/T06.'),
 ('Backup và chỉ huy DR','Bản sao, phụ thuộc phục hồi, failover và failback.','BIA, job, runbook, T08/T09/T10.'),
 ('IDC và mua sắm','Vật lý, nguồn điện, PCCC, ra vào và nhà cung cấp.','Hồ sơ từng site, SLA, phân định trách nhiệm.'),
 ('Người kiểm tra','Kiểm tra phạm vi, bằng chứng, kết quả và retest.','Biên bản, kết luận có giới hạn, danh sách tồn tại.')],[1.1,2.2,2.1])
d.p('Dịch vụ khách hàng tự quản phải có ma trận trách nhiệm chung. Hạ tầng nền đạt yêu cầu không tự chứng minh ứng dụng, hệ điều hành khách, dữ liệu hoặc tài khoản của khách hàng đã được phê duyệt cùng cấp độ.')
d.page('3 Thu thập và quản lý bằng chứng')
d.p('Mỗi bằng chứng dùng mã EV và liên kết tới mã TC, CL, HS hoặc T. Tên tệp gợi ý: EV-001_TC-001_DC_20261017_v1. Ghi ngày giờ, phạm vi, bản dựng, người thu, người xác minh và đường dẫn kho. Tạo bản đã che dữ liệu để đưa vào hồ sơ; giữ bản gốc ở kho có phân quyền và kiểm soát toàn vẹn.')
d.table(['Nguồn','Cách thu','Lưu ý kiểm tra'],[
 ('XOR STACK CLOUD','Xuất inventory service, endpoint, project/role, network, port, SG, image và volume theo quyền được cấp.','Ghi phiên bản CLI/API và project scope; không xuất token hoặc mật khẩu.'),
 ('Host/cluster','Thu phiên bản OS/gói, cấu hình đã lọc bí mật, cluster status, route và trạng thái agent.','Bản cấu hình trích phải gắn host, ngày, hash; không dùng ảnh menu làm chứng cứ tính năng đang hoạt động.'),
 ('IAM/PAM/KMS','Xuất chính sách, vai trò, log mẫu, lịch sử cấp/thu hồi và luân chuyển.','Không đưa khóa riêng, recovery code hoặc secret vào tài liệu.'),
 ('SOC/backup','Xuất nguồn log, rule, retention, job, mẫu cảnh báo và biên bản restore.','Có số đo thời gian và bản ghi xử lý; ảnh dashboard riêng lẻ chưa đủ.'),
 ('IDC','Thu hồ sơ nghiệm thu, bảo trì, điện, camera, ra vào và hợp đồng.','Xác minh phạm vi chứng chỉ bao phủ đúng cơ sở/phòng/rack và còn hiệu lực.')],[1,2.2,2])
d.h('Thứ tự ưu tiên bằng chứng')
d.p('Kết hợp tài liệu đã phê duyệt, cấu hình thực tế, log vận hành và phép thử. Bằng chứng phải trả lời được yêu cầu đã triển khai ở đâu, có hoạt động hay không, ai chịu trách nhiệm và còn hiệu lực tại thời điểm đánh giá. Không thay phép thử bằng cam kết mua sắm hoặc mô tả tính năng sản phẩm.')
d.page('4 Các mốc kiểm tra định kỳ')
d.table(['Hoạt động','Mốc tối thiểu theo TCVN','Điều khoản'],[
 ('Nhận diện và xem lại hiệu lực xử lý rủi ro','1 lần/6 tháng và khi có thay đổi liên quan','6.1.2.2; 6.1.2.4'),
 ('Rà thiết bị trái phép, tài sản kết nối, DHCP/IPAM','1 lần/quý','6.2.2.2–6.2.2.4'),
 ('Danh mục tài khoản và quyền dữ liệu','1 lần/quý','6.6.2.1; 6.4.2.1'),
 ('Scan toàn hệ thống / tài sản quan trọng','1 lần/6 tháng / 1 lần/quý','6.7.2.1'),
 ('Máy người dùng và di động được cấp','Kiểm tra/cập nhật vá ít nhất 1 lần/tháng nếu có','6.7.2.2'),
 ('Rà nhật ký an ninh','1 lần/tháng','6.8.2.1'),
 ('Lưu nhật ký an ninh / camera phòng máy','Tối thiểu 6 tháng cho mỗi loại','6.8.2.1; A.4.2'),
 ('Thử phục hồi bản sao','1 lần/6 tháng','6.11.2.5'),
 ('Rà sơ đồ mạng','1 lần/6 tháng và khi thay đổi','6.12.2.1'),
 ('Điều chỉnh ngưỡng cảnh báo','1 lần/quý','6.13.2.5'),
 ('Đào tạo và diễn tập sự cố','1 lần/năm; theo vai trò và thay đổi khi áp dụng','6.14.2; 6.16.2.6; 6.17.2.6'),
 ('Danh mục thư viện/thành phần bên thứ ba','1 lần/quý','6.16.2.4'),
 ('Pentest bên trong và bên ngoài','Mỗi loại ít nhất 1 lần/năm','6.18.2.2; 6.18.2.5')],[2,2,1.3],10)
d.p('Các quy trình và danh mục còn có mốc rà soát riêng tại từng dòng TCVN. RPO/RTO, thời gian phản ứng SOC, thời hạn khắc phục lỗ hổng và số ngày bất biến phải được XOR phê duyệt theo rủi ro/hợp đồng; không gán một giá trị mẫu thành yêu cầu pháp luật.')
for k in range(0,len(TESTS),2):
 d.page(f'{5+k//2} Kịch bản kiểm thử {TESTS[k][0]} và {TESTS[k+1][0]}')
 for x in TESTS[k:k+2]:
  d.h(x[0]+' '+x[1]);d.p('Đầu mối: '+x[2]+'. Căn cứ: '+x[3]+'.')
  d.p('Điều kiện trước thử: '+x[4]);d.p('Thực hiện: '+x[5]);d.p('Tiêu chí chấp nhận: '+x[6]);d.p('Lưu bằng chứng: '+x[7])
d.page('11 Biên bản và điều kiện chuyển hồ sơ')
d.table(['Trường biên bản','Nội dung cần ghi'],[
 ('Mã và phạm vi','Mã T/TC/CL; site; asset ID; project; bản dựng; ngày giờ; cửa sổ thử.'),
 ('Thành phần tham gia','Người thực hiện, người chứng kiến, người xác nhận và phạm vi quyền.'),
 ('Kết quả','Điều kiện đầu vào, bước đã làm, số đo, kết quả so với tiêu chí và mã EV.'),
 ('Tồn tại','Mô tả, ảnh hưởng, mức ưu tiên theo rủi ro, đầu mối và hạn.'),
 ('Retest','Ngày thử lại, phiên bản mới, kết quả, bằng chứng và người xác nhận.'),
 ('Kết luận','Đạt/Không đạt/Chưa đủ bằng chứng trong phạm vi cụ thể; các giới hạn và ngoại lệ.')],[1,3])
d.p('Đợt rà soát được đóng khi phạm vi và danh mục đã chốt, các mục áp dụng có bằng chứng được xác minh, kết quả thử được lưu, phát hiện được xử lý theo quyết định đúng thẩm quyền. Việc chấp nhận rủi ro không tự thay thế một yêu cầu bắt buộc chưa đáp ứng.')
d.p('Trước gửi thẩm định: chuyển dữ liệu thực tế sang Bộ 2; đính kèm thiết kế đã phê duyệt hoặc tài liệu tương đương; hoàn thiện đánh giá rủi ro chi tiết và ý kiến chuyên trách. Không gửi bản còn [CẦN ĐIỀN], không dùng bảng tổng hợp Excel như quyết định phê duyệt.')
sources(d,True);d.topology('Phụ lục');d.save(TECH/'01_So_tay_ra_soat_ky_thuat.docx')

# DANH MỤC HỒ SƠ
d=Doc('Danh mục hồ sơ đề xuất cấp độ 4')
d.p('CÔNG TY CỔ PHẦN XOR CLOUD',True,align=WD_ALIGN_PARAGRAPH.CENTER)
d.p('DANH MỤC HỒ SƠ',True,align=WD_ALIGN_PARAGRAPH.CENTER)
d.p('Đề xuất cấp độ 4 hệ thống thông tin XOR SECURE CLOUD',True,align=WD_ALIGN_PARAGRAPH.CENTER)
d.p('Chủ quản: '+ORG+'. Đơn vị vận hành: '+OPERATOR+'.')
d.p('Mã hồ sơ: XOR-CD4. Phiên bản: 2.4. Ngày lập: [CẦN ĐIỀN].')
files=[('01','Thuyết minh hồ sơ đề xuất cấp độ 4'),('02','Phụ lục đối chiếu yêu cầu TCVN'),('03','Quy chế và quy trình bảo đảm an ninh mạng'),('04','Công văn xin ý kiến chuyên môn'),('05','Công văn đề nghị thẩm định'),('06','Tờ trình phê duyệt cấp độ'),('07','Quyết định phê duyệt cấp độ'),('08','Ý kiến chuyên môn'),('09','Báo cáo kết quả rà soát'),('10','Phụ lục thiết kế và hiện trạng'),('11','Công văn đề nghị đánh giá độc lập'),('12','Thông báo về cấp độ hệ thống thông tin')]
d.table(['STT','Tên tài liệu','Số / ngày văn bản'],[(n,t,'[CẦN ĐIỀN]') for n,t in files],[.45,2.9,1.1])
d.p('Tài liệu kèm theo: [CẦN ĐIỀN danh mục thiết kế được phê duyệt, ý kiến thẩm định, bằng chứng đánh giá và tài liệu pháp lý; số, ngày, phiên bản, số trang].')
d.p('Đầu mối: '+AUTHOR_NAME+' — '+AUTHOR_ROLE+'. Liên hệ: [CẦN ĐIỀN điện thoại, email].')
d.save(ADMIN/'00_Danh_muc_ho_so.docx')

# THUYẾT MINH CHÍNH
d=Doc('Hồ sơ đề xuất cấp độ 4 hệ thống XOR SECURE CLOUD')
d.cover('Tài liệu thuyết minh theo Điều 21 và Điều 22 Nghị định 331 năm 2026')
d.table(['Thông tin quản lý hồ sơ','Nội dung'],[
 ('Chủ quản',ORG),('Đơn vị vận hành',OPERATOR),('Địa chỉ trụ sở',ADDRESS),('Đại diện chủ quản',SIGNER+' — Tổng Giám đốc'),('Mã hồ sơ','XOR-CD4-HS-01 — phiên bản 2.4'),('Tình trạng','Lập lần đầu; chưa có quyết định phê duyệt cấp độ hoặc chứng nhận.'),('Phạm vi','Toàn bộ XOR STACK CLOUD và dịch vụ quản trị chung tại DC Cầu Giấy, DR Nam Thăng Long.')],[1.2,2.8])
d.p('Hồ sơ đề xuất cấp độ 4 đối với toàn bộ nền tảng XOR SECURE CLOUD và các dịch vụ quản trị chung, bao gồm thuyết minh hệ thống, phương án thiết kế, đánh giá rủi ro và phương án bảo đảm an ninh mạng.')
d.page('Phần I Tổng quan về hệ thống thông tin')
d.h('1 Căn cứ lập hồ sơ')
d.p('Căn cứ Luật An ninh mạng số 116/2025/QH15; '+LEGAL+'; căn cứ chức năng, nhiệm vụ và thẩm quyền của chủ quản [CẦN ĐÍNH KÈM VĂN BẢN]; căn cứ hiện trạng và nhu cầu bảo vệ hệ thống XOR SECURE CLOUD, đơn vị vận hành lập hồ sơ đề xuất cấp độ 4 để xin ý kiến chuyên môn, thẩm định và trình phê duyệt theo quy định.')
d.p('Phương án bảo đảm an ninh mạng được đối chiếu với Điều 6 và Phụ lục A.4 TCVN 14423:2026. [S1, Điều 21–22; S2]')
d.h('2 Thông tin chủ quản và đơn vị vận hành')
d.table(['Trường thông tin','Chủ quản','Đơn vị vận hành'],[
 ('Tên đầy đủ',ORG,OPERATOR),('Mã số doanh nghiệp/định danh','[CẦN ĐIỀN theo đăng ký doanh nghiệp]','Đơn vị nội bộ của chủ quản'),('Chức năng, nhiệm vụ và quyền hạn','[CẦN ĐIỀN; dẫn văn bản]','[CẦN ĐIỀN; dẫn quyết định giao vận hành]'),('Đại diện và đầu mối phụ trách',SIGNER+' — Tổng Giám đốc',OPS_SIGNER+' — Giám đốc Khối Hạ tầng & Vận hành'),('Địa chỉ và liên hệ',ADDRESS+'; [điện thoại/email]', ADDRESS+'; [CẦN ĐIỀN điện thoại/email]'),('Đơn vị chuyên trách ANM',SPECIALIST+'; [CẦN ĐÍNH KÈM quyết định giao nhiệm vụ]',OPS_SIGNER+'; [CẦN XÁC NHẬN phân công và cơ chế độc lập]')],[1,1.5,1.5],11)
d.p('Đầu mối lập hồ sơ: '+AUTHOR_NAME+' — '+AUTHOR_ROLE+'. [CẦN ĐIỀN số, ngày văn bản xác định tư cách chủ quản, giao nhiệm vụ vận hành/chuyên trách, thẩm quyền ký và cơ chế kiểm tra độc lập].')
d.page('3 Chức năng phạm vi và quy mô')
d.p('XOR SECURE CLOUD là hệ thống cung cấp và quản lý tài nguyên điện toán đám mây trên nền tảng XOR STACK CLOUD. Phạm vi hồ sơ gồm hạ tầng, phần mềm nền tảng và các dịch vụ quản trị chung tại Trung tâm Hạ tầng IDC Cầu Giấy (DC) và Trung tâm Hạ tầng IDC Nam Thăng Long (DR).')
d.p('Phạm vi kỹ thuật bao gồm máy chủ, ảo hóa, mạng, lưu trữ, portal/API, control plane, định danh và đặc quyền, quản lý khóa/chứng thư, DNS/thời gian, giám sát/nhật ký, sao lưu/DR, chuỗi phát hành và máy trạm quản trị. Danh mục dịch vụ và các thành phần thuê ngoài được xác định tại Phụ lục PL01. [S1, Điều 22.3]')
d.table(['Nội dung','Thông tin và tài liệu chứng minh'],[
 ('Đối tượng phục vụ','[CẦN ĐIỀN nhóm khách hàng, địa bàn, số tổ chức/tenant/người dùng và phạm vi toàn quốc nếu có].'),
 ('Quy mô hai IDC','[CẦN ĐIỀN số rack, host/controller/compute/storage/network; dung lượng; tải và năng lực dự phòng từng site].'),
 ('Dịch vụ','[CẦN CHỐT IaaS/VM, network, block/object/file, K8s, DBaaS, GPU và dịch vụ quản lý thực tế].'),
 ('Thông tin xử lý','[CẦN ĐIỀN loại dữ liệu, độ nhạy cảm, dữ liệu cá nhân/quan trọng và dữ liệu thuộc quy định chuyên ngành nếu có].'),
 ('Kết nối ngoài','[CẦN ĐIỀN ISP, hệ thống khách hàng, đối tác, dịch vụ dùng chung, cơ quan quản lý và kênh liên site].'),
 ('SLA và thời gian hoạt động','[CẦN ĐIỀN nghĩa vụ 24/7, thời gian bảo trì, RPO/RTO và điều khoản hợp đồng/BIA].')],[1.1,3.1])
d.h('4 Ranh giới trách nhiệm')
d.p('XOR chịu trách nhiệm đối với hạ tầng, nền tảng và dịch vụ đã cam kết quản lý. Chủ quản của từng hệ thống khách hàng chịu trách nhiệm mục đích sử dụng, phân loại dữ liệu, tài khoản và các cấu hình/ứng dụng thuộc phần mình quản lý. Ranh giới phải được cụ thể hóa theo hợp đồng, dịch vụ và ma trận trách nhiệm; việc phê duyệt cấp độ nền tảng không thay thế hồ sơ cấp độ của từng hệ thống khách hàng.')
d.page('Phần II Thiết kế và hiện trạng hệ thống')
d.p('Thiết kế hệ thống gồm kiến trúc tổng thể DC–DR, phân vùng mạng, các luồng kết nối, hạ tầng vật lý và các thành phần quản trị. Hồ sơ thiết kế được phê duyệt hoặc tài liệu có giá trị tương đương: [CẦN ĐIỀN tên, số, ngày, cơ quan phê duyệt và danh mục bản vẽ]. [S1, Điều 21.2 b]')
d.h('1 Danh mục thiết kế và phụ lục hiện trạng')
d.table(['Mã','Tài liệu phải kèm theo','Nội dung'],[
 ('PL01','Dịch vụ và ứng dụng','Tên, chức năng, đối tượng, máy chủ/vị trí, OS, phiên bản, chủ sở hữu và phạm vi XOR quản lý.'),
 ('PL02','Thiết bị và cấu hình vật lý','ID, loại, model/serial, rack, site, vai trò, kết nối, nguồn, phiên bản và EOL.'),
 ('PL03','Phân vùng mạng và IP','VLAN/VRF/VXLAN, CIDR, IP private/public, route, endpoint, owner và mục đích.'),
 ('PL04','Luồng dữ liệu và kết nối','Nguồn/đích/cổng/giao thức, phân loại dữ liệu, xác thực, mã hóa, log, tần suất và phê duyệt.'),
 ('PL05','Danh mục dữ liệu và quyền','Chủ dữ liệu, mức nhạy cảm, kho, quyền, thời hạn lưu, hủy, sao lưu và khóa.'),
 ('PL06','HA backup và DR','Phụ thuộc, quorum, miền lỗi, capacity, replica, RPO/RTO, runbook, restore và diễn tập.'),
 ('PL07','Cấu hình và phiên bản','Baseline, inventory, SBOM, khác biệt XOR/upstream, CI/CD, phê duyệt và rollback.')],[.5,1.3,2.7],11)
d.p('Hiện trạng triển khai: [CẦN ĐIỀN số, ngày hồ sơ kiểm kê/nghiệm thu; phiên bản cấu hình; sai khác giữa thiết kế và hiện trạng; mã tài liệu tại PL01–PL07].')
d.topology()
d.h('3 Phương án phân vùng và luồng kết nối',1)
d.table(['Vùng','Chức năng và giới hạn'],[
 ('Z01 Biên','Kết nối ISP, chống DDoS, NGFW/IPS và cổng VPN/ZTNA; tách listener quản trị.'),
 ('Z02 DMZ','WAF, API gateway/LB và endpoint công khai được duyệt; chặn truy cập trực tiếp origin.'),
 ('Z03 Nội bộ','Người dùng nội bộ và thiết bị; kiểm soát kết nối bằng NAC hoặc tương đương.'),
 ('Z04 Tài nguyên tenant','VM/container, project/SG/policy; tách tenant và môi trường.'),
 ('Z05 Dữ liệu','DB, block/object/file; tách mạng lưu trữ, quyền và khóa theo phạm vi.'),
 ('Z06 Quản trị','PAM, IAM, PKI/KMS và dịch vụ quản trị; phân khu chức năng; không truy cập Internet trực tiếp.'),
 ('Z07 Điều khiển','Controller/API nội bộ, DB trạng thái và hàng đợi; mTLS/ACL theo phiên bản.'),
 ('Z08 An ninh','SOC, collector, SIEM và kho log khác vùng phát sinh; quyền quản trị độc lập.'),
 ('Z09 Sao lưu','Kho và miền quản trị độc lập; kiểm soát xóa, phiên bản và mã hóa.'),
 ('Z10 Ngoài băng','BMC/console và quản trị thiết bị; chỉ qua đích quản trị được phép.'),
 ('Z11 Phát triển','Git/CI/build/registry/dev/test, tách production và bí mật.'),
 ('Z12 Liên site','Đồng bộ DC–DR trên kết nối riêng được bảo vệ; chỉ cho phép peer/luồng đã duyệt.')],[1.2,3.4],11)
d.p('Mọi vùng áp dụng chặn mặc định; quyền kết nối được xác định theo nhu cầu dịch vụ và có người sở hữu. Quyền quản trị đi qua cổng kiểm soát và PAM; không cấp tuyến mạng rộng cho người truy cập từ xa. Phương án kiểm soát ranh giới lõi với vùng ít tin cậy phải thuyết minh Datadiode/USG/BSG hoặc biện pháp tương đương phù hợp từng luồng. [S2, 6.12.2.1–6.12.2.6]')
d.page('4 Phương án nền tảng dữ liệu và phục hồi')
d.h('4 1 Nền tảng XOR STACK CLOUD')
d.p('Mỗi IDC được thiết kế thành cụm độc lập, có quorum nội bộ cho các thành phần yêu cầu đồng thuận. Việc bố trí số controller, DB, queue, compute và storage căn cứ vào bản dựng được hỗ trợ và phép đo tải. Cụm đồng thuận được thiết lập trong từng IDC; rủi ro miền lỗi chung giữa hai IDC được đánh giá trong phương án DR. Hệ thống mạng ảo hóa/SDN, hypervisor và storage được kiểm tra cách ly tenant, metadata, console, snapshot và tái cấp phát tài nguyên.')
d.h('4 2 Bảo vệ dữ liệu và khóa')
d.p('Dữ liệu được phân loại theo chủ sở hữu, độ nhạy cảm và mục đích xử lý. Quyền đọc, ghi, xuất, sao lưu và hủy được kiểm soát. Phương án mã hóa dữ liệu quan trọng khi truyền/lưu, bảo vệ khóa, toàn vẹn, chữ ký số, chống thất thoát và nhật ký truy cập được thuyết minh theo từng dòng 6.4 tại tài liệu 02. Giải pháp kiểm soát tương đương được thuyết minh và kiểm chứng theo phạm vi áp dụng.')
d.h('4 3 Sao lưu và DR')
d.p('Bản sao dữ liệu, cấu hình và thông tin phục hồi phải được lưu trên hạ tầng độc lập với môi trường vận hành; có phiên bản, mã hóa phù hợp, quản lý quyền và phép thử restore. Danh mục sao lưu bao gồm control plane, dữ liệu dịch vụ thuộc trách nhiệm XOR, cấu hình mạng, image và khóa/chứng thư cần khôi phục. RPO, RTO, dung lượng DR và lịch đồng bộ được chốt theo BIA tại Phần III.')
d.p('Kịch bản mất DC phải xác nhận nguồn dữ liệu hợp lệ, cô lập nguồn ghi cũ trước kích hoạt DR, phục hồi cả IAM/DNS/KMS và dịch vụ phụ thuộc, đối soát nghiệp vụ và thực hiện chuyển về có kiểm soát. Hai IDC trong cùng khu vực phải được đánh giá rủi ro đồng thời về điện, tuyến cáp, ISP và thiên tai; có biện pháp đối với miền lỗi chung.')
d.h('4 4 Độc lập dịch vụ phục hồi')
d.p('Tài liệu khôi phục và bí mật khẩn cấp được lưu ở nơi truy cập được khi nền tảng dừng. Đơn vị vận hành lập đồ thị phụ thuộc và kiểm thử khởi động nguội để phát hiện vòng phụ thuộc. [CẦN ĐIỀN mã bằng chứng triển khai và biên bản kiểm thử T07–T10].')
d.page('Phần III Thuyết minh đề xuất cấp độ và rủi ro')
d.h('1 Loại hệ thống và tiêu chí đề xuất')
d.p('XOR SECURE CLOUD thuộc loại hệ thống cơ sở hạ tầng thông tin cung cấp tài nguyên điện toán đám mây. Cấp độ đề xuất: cấp độ 4. Chức năng, phạm vi phục vụ, loại thông tin và hậu quả rủi ro được thuyết minh theo Điều 14 và Điều 22 Nghị định 331/2026/NĐ-CP.')
d.table(['Tiêu chí Điều 14','Nội dung cần chứng minh đối với XOR','Tình trạng'],[
 ('Khoản 1','Thông tin bí mật nhà nước/quốc phòng/an ninh và hậu quả nghiêm trọng [CẦN ĐIỀN phạm vi phục vụ nếu áp dụng].','[CẦN XÁC NHẬN áp dụng và mã bằng chứng]'),
 ('Khoản 2','Hạ tầng phục vụ cơ quan/tổ chức phạm vi toàn quốc, yêu cầu 24/7 và không chấp nhận ngừng không có kế hoạch. Phạm vi khách hàng, nghĩa vụ dịch vụ và BIA: [CẦN ĐIỀN].','[CẦN XÁC NHẬN áp dụng và mã bằng chứng]'),
 ('Khoản 3','Hệ thống điều khiển công nghiệp phục vụ công trình thuộc trường hợp được quy định.','[CẦN XÁC NHẬN áp dụng và mã bằng chứng]'),
 ('Khoản 4','Quyết định của Thủ tướng hoặc kết quả đánh giá theo Khung quản lý rủi ro an ninh mạng trong trường hợp tương ứng.','[CẦN ĐIỀN văn bản làm căn cứ nếu áp dụng]')],[.7,2.9,1.1],11)
d.p('Tiêu chí đề xuất cấp độ: [CẦN ĐIỀN khoản áp dụng tại Điều 14]. Căn cứ và tài liệu chứng minh: [CẦN ĐIỀN phạm vi phục vụ, hợp đồng, BIA, kết quả đánh giá rủi ro và mã tài liệu].')
d.h('2 Hệ thống liên quan')
d.p('[CẦN LẬP danh mục hệ thống kết nối, phụ thuộc hoặc bị ảnh hưởng khi XOR dừng; tên chủ quản, cấp độ đã được phê duyệt nếu có, dịch vụ trao đổi, loại dữ liệu và mức tác động]. Các phụ thuộc IAM, DNS, khóa, ISP, IDC và nhà cung cấp được xác định tại PL04 và BIA. [S1, Điều 22.5 a]')
d.page('3 Báo cáo đánh giá rủi ro chi tiết')
d.p('Phương pháp đánh giá rủi ro gồm xác định tài sản và dịch vụ; nhận diện mối đe dọa, điểm yếu; xác định khả năng và tác động; đánh giá hiệu lực kiểm soát; lựa chọn biện pháp xử lý; đánh giá rủi ro còn lại và phê duyệt. [CẦN ĐIỀN phương pháp, thang điểm và văn bản phê duyệt áp dụng]. [S1, Điều 22.5; S2, 6.1]')
d.table(['Rủi ro cần đánh giá','Tác động phải làm rõ','Kiểm soát dự kiến'],[(x[0]+' '+x[1],x[3],x[2]) for x in RISKS],[1.5,1.5,1.6],11)
d.p('Kết quả đánh giá rủi ro: [CẦN ĐIỀN tài sản, mức khả năng, mức tác động, căn cứ chấm điểm, biện pháp xử lý, rủi ro còn lại, người chịu trách nhiệm, thời hạn và quyết định chấp nhận].')
d.page('4 Phân tích tác động và yêu cầu liên tục')
d.p('Phân tích tác động nghiệp vụ xác định số khách hàng và hệ thống bị ảnh hưởng; nguy cơ mất, lộ, sửa dữ liệu; thời gian gián đoạn; hậu quả tài chính, pháp lý và tác động đối với lợi ích công cộng, trật tự, an toàn xã hội, an ninh quốc gia. [CẦN ĐIỀN báo cáo BIA, số liệu, nguồn và người xác nhận]. [S1, Điều 22.5 b c]')
d.table(['Dịch vụ','Phụ thuộc cần đối chiếu','Chỉ tiêu và căn cứ'],[(x[0]+' '+x[1],x[2],'[CẦN ĐIỀN RPO, RTO, MTD; hợp đồng/BIA; người phê duyệt]') for x in SERVICES],[1.2,1.9,1.7],11)
d.p('Yêu cầu hoạt động liên tục: [CẦN ĐIỀN nghĩa vụ 24/7; điều kiện về ngừng vận hành có kế hoạch và không có kế hoạch; bên yêu cầu, phạm vi dịch vụ, tác động và biện pháp đáp ứng]. [S1, Điều 22.5 d]')
d.page('Phần IV Phương án bảo đảm an ninh mạng')
d.h('1 Biện pháp quản lý và kỹ thuật')
d.p('Phương án đối với từng yêu cầu TCVN được xác định tại Phụ lục đối chiếu yêu cầu và phương án bảo đảm an ninh mạng. Quy chế và quy trình quy định trách nhiệm, trình tự thực hiện và hồ sơ lưu. Hiện trạng và kết quả kiểm tra: [CẦN ĐIỀN ngày chốt, danh mục cấu hình, văn bản ban hành, biên bản kiểm tra và mã bằng chứng]. [S1, Điều 22.6]')
d.table(['Nhóm','Phương án tổ chức thực hiện'],[
 ('Quản trị rủi ro và tài sản','Lập CMDB, phân loại dữ liệu, BIA, quản lý thay đổi và sổ rủi ro; cập nhật theo tần suất từng điều khoản.'),
 ('Định danh và đặc quyền','Tài khoản cá nhân, quyền tối thiểu, MFA, PAM, tách vận hành/bảo mật/kiểm toán, thu hồi theo vòng đời.'),
 ('Mạng và nền tảng','Phân vùng, microsegmentation, HA, FW/IPS/WAF/DDoS/DBF/NAC hoặc giải pháp tương đương có kiểm chứng; bảo vệ API và tenant.'),
 ('Dữ liệu và khôi phục','Mã hóa, khóa, toàn vẹn, DLP, chữ ký số khi áp dụng; sao lưu độc lập; kiểm thử restore và DR.'),
 ('Giám sát và ứng phó','SOC 24/7, log tối thiểu 6 tháng, đồng bộ thời gian, kho tách mạng, EDR, cảnh báo và playbook.'),
 ('Phát triển và đánh giá','Phát triển an toàn phần mềm nền tảng XOR; SBOM, kiểm tra mã nguồn/dependency, kiểm thử trước phát hành, pentest trong/ngoài, retest.')],[1.1,3.1])
d.h('2 Điều kiện đối với giải pháp tương đương')
d.p('Giải pháp tương đương được đánh giá theo yêu cầu áp dụng, rủi ro, kiến trúc kiểm soát, khả năng bảo vệ, giới hạn, kết quả thử nghiệm và thẩm quyền chấp thuận. Hồ sơ giải pháp tương đương: [CẦN ĐIỀN mã yêu cầu, thuyết minh, kết quả kiểm chứng và văn bản chấp thuận nếu áp dụng].')
d.page('3 Bảo đảm an toàn vật lý và trách nhiệm thuê ngoài')
d.p('Phương án bảo đảm an toàn vật lý áp dụng riêng tại DC Cầu Giấy và DR Nam Thăng Long theo Phụ lục A.4. Phạm vi kiểm soát gồm phòng máy, rack thiết bị, điện/UPS, phòng cháy và chữa cháy, chống sét, chống nước, tĩnh điện, môi trường, bảo vệ điện từ, camera và kiểm soát ra vào. Hồ sơ đối chiếu của từng IDC: [CẦN ĐIỀN hợp đồng, phạm vi, thời hạn hiệu lực, biên bản kiểm tra và mã bằng chứng].')
d.table(['Nội dung','XOR','IDC hoặc bên cung cấp'],[
 ('Rack thiết bị và tài sản','Quản lý thiết bị, người được phép, niêm phong/bàn giao và inventory.','Hạ tầng phòng máy, điều kiện vật lý và kiểm soát vào/ra theo hợp đồng.'),
 ('Điện và môi trường','Đấu nối, giám sát tải, cảnh báo và phối hợp bảo trì.','Bằng chứng hai nguồn từ hai trạm khác nhau, UPS, điều hòa, chống cháy/nước/sét theo A.4.'),
 ('Camera và ra vào','Phê duyệt người, rà soát log liên quan, xử lý bất thường.','Giám sát 24/7, lưu camera tối thiểu 6 tháng, cung cấp trích xuất theo thẩm quyền.'),
 ('Sự cố','Chỉ huy phản ứng với hệ thống, thông tin khách hàng theo trách nhiệm.','Thông báo sự cố cơ sở, hỗ trợ điều tra/phục hồi theo SLA.'),
 ('Kiểm tra định kỳ','Tổ chức kiểm tra, đối soát bằng chứng và theo dõi khắc phục.','Cung cấp hồ sơ, tham gia kiểm tra theo hợp đồng và phạm vi cho phép.')],[1,1.8,2])
d.h('4 Ranh giới trách nhiệm với khách hàng')
d.p('Phụ lục hợp đồng xác định bên quản lý OS, ứng dụng, DB, dữ liệu, khóa, tài khoản, sao lưu và log. Đối với IaaS, khách hàng thường quản lý OS khách và ứng dụng; với dịch vụ quản lý, phần XOR chịu trách nhiệm phải được mô tả cụ thể. Cấp độ được công bố chỉ gắn với phạm vi đã phê duyệt và không chuyển toàn bộ trách nhiệm của chủ quản khách hàng sang XOR.')
d.page('Phần V Kế hoạch hoàn thiện và tài liệu kèm theo')
d.table(['Giai đoạn','Kết quả phải có','Đầu mối và hạn'],[
 ('1 Xác minh đầu vào','Pháp nhân, thẩm quyền, phạm vi, tiêu chí Điều 14, dữ liệu và dịch vụ.','[CẦN ĐIỀN người và hạn]'),
 ('2 Xác nhận hiện trạng','Bộ as-built, CMDB, phiên bản, IP và các phụ lục PL01–PL07.','[CẦN ĐIỀN]'),
 ('3 Bổ sung kiểm soát','Theo từng thiếu sót trong TCVN; phê duyệt thay đổi, triển khai và lưu bằng chứng.','[CẦN ĐIỀN]'),
 ('4 Kiểm tra và retest','T01–T12; báo cáo độc lập, BIA/RPO/RTO đo được, đóng phát hiện.','[CẦN ĐIỀN]'),
 ('5 Ý kiến và thẩm định','Ý kiến chuyên trách; công văn và bộ hồ sơ đúng phiên bản; giải trình khi yêu cầu.','[CẦN ĐIỀN]'),
 ('6 Phê duyệt và duy trì','Quyết định chủ quản; thông báo đúng phạm vi; lịch vận hành, kiểm tra và cập nhật.','[CẦN ĐIỀN]')],[.8,2.7,1.1])
d.page('Danh mục kèm theo hồ sơ trình')
for x in DOSSIER[:9]:d.p(x[0]+' '+x[1]+': [CẦN ĐIỀN tên tệp, phiên bản, số/ngày và số trang thực tế].')
sources(d)
annexes=[
 ('PL01','Danh mục dịch vụ và ứng dụng',[
  ('Mã dịch vụ; tên; chức năng','[CẦN ĐIỀN]'),('Đối tượng và phạm vi phục vụ','[Nhóm khách hàng; địa bàn; số tenant; nghiệp vụ]'),
  ('Chủ dịch vụ; đơn vị quản lý','[Tên; chức năng; đầu mối]'),('Thành phần và vị trí triển khai','[DC/DR; asset ID; cụm; vùng mạng; endpoint]'),
  ('Phần mềm và hệ điều hành','[Tên; phiên bản; bản dựng; giấy phép; EOL]'),('Dữ liệu và dịch vụ phụ thuộc','[Mã PL05; IAM/DNS/KMS/DB/queue/storage và bên ngoài]'),
  ('Ranh giới trách nhiệm','[XOR/khách hàng/nhà cung cấp; dẫn hợp đồng]'),('Năng lực và nghĩa vụ dịch vụ','[Tải; dung lượng; SLA; lịch hoạt động; dẫn BIA]'),
  ('Tài liệu xác nhận','[Hợp đồng; CMDB; thiết kế; biên bản; mã EV]')]),
 ('PL02','Thiết bị và cấu hình vật lý',[
  ('Mã tài sản; loại; vai trò','[CẦN ĐIỀN mã tài sản, loại và vai trò]'),('Site và vị trí','[Cầu Giấy/Nam Thăng Long; phòng; rack; U; địa chỉ thực tế]'),
  ('Nhà sản xuất; model; serial','[CẦN ĐIỀN]'),('Cấu hình và phiên bản','[CPU/RAM/disk/NIC; firmware; hypervisor/OS; license]'),
  ('Kết nối vật lý','[Cổng; peer; fabric A/B; tốc độ; loại cáp; bản vẽ số]'),('Nguồn và miền lỗi','[PDU/UPS/nguồn A/B; rack; failure domain]'),
  ('Bảo trì và vòng đời','[Đơn vị; hợp đồng; EOL/EOS; lịch bảo trì]'),('Trách nhiệm và hiện trạng','[Chủ sở hữu; người quản lý; hoạt động/dự phòng; ngày kiểm kê]'),
  ('Hồ sơ hiện trạng','[Bản vẽ vật lý; số/ngày phê duyệt hoặc tài liệu tương đương; EV]')]),
 ('PL03','Phân vùng mạng và địa chỉ',[
  ('Mã mạng; chức năng; site','[CẦN ĐIỀN; vùng Z01–Z12 tương ứng]'),('Định danh phân đoạn','[VLAN; VRF; VXLAN/VNI; project; network ID]'),
  ('Địa chỉ và điểm cuối','[CIDR; gateway; IP private/public; VIP; FQDN; asset ID]'),('Định tuyến','[Bảng route; next-hop; miền định tuyến; đích được phép]'),
  ('Ranh giới và kiểm soát','[FW/SG/ACL/SDN; mã rule; chặn mặc định; chống giả mạo]'),('Quản trị và truy cập Internet','[Đường PAM; OOB; kiểm chứng mạng quản trị không kết nối Internet]'),
  ('HA và phụ thuộc','[Peer; fabric; LB; DNS; NTP; khả năng failover]'),('Chủ mạng và thay đổi','[Đơn vị; đầu mối; ticket; ngày phê duyệt]'),
  ('Hồ sơ hiện trạng','[Bản vẽ lô-gic; bảng IP; export route/rule đã xử lý bí mật; EV]')]),
 ('PL04','Luồng dữ liệu và kết nối hệ thống',[
  ('Mã luồng; mục đích','[CẦN ĐIỀN; căn cứ nghiệp vụ]'),('Hệ thống nguồn và đích','[Tên; chủ quản; cấp độ đã phê duyệt nếu có; bên chịu trách nhiệm]'),
  ('Điểm kết nối','[Site; vùng; asset/IP/FQDN; cổng; giao thức; chiều khởi tạo]'),('Thông tin trao đổi','[Mã dữ liệu PL05; phân loại; dung lượng/tần suất]'),
  ('Xác thực và quyền','[Danh tính kỹ thuật; phạm vi quyền; cơ chế xác thực; chu kỳ đổi khóa]'),('Mã hóa và kiểm tra toàn vẹn','[Giao thức; bộ mật mã; chứng thư; cơ chế bảo vệ theo phân loại]'),
  ('Kiểm soát tại biên','[Rule/allowlist; WAF/IPS/gateway; kiểm soát tương đương nếu áp dụng]'),('Nhật ký và giám sát','[Trường log; đích collector; retention; cảnh báo; người quản lý]'),
  ('Phê duyệt và kiểm thử','[Người duyệt; thời hạn; ticket; phép thử cho phép/từ chối; EV]')]),
 ('PL05','Dữ liệu và quyền truy cập',[
  ('Mã tập dữ liệu; mô tả','[CẦN ĐIỀN; nguồn và mục đích xử lý]'),('Chủ sở hữu và trách nhiệm','[Chủ dữ liệu; bên xử lý; XOR/khách hàng; căn cứ hợp đồng]'),
  ('Phân loại và căn cứ','[Độ nhạy cảm; nghĩa vụ pháp lý/chuyên ngành; người phê duyệt]'),('Nơi lưu và bản sao','[Site; hệ thống; volume/bucket/DB; replica; backup; môi trường thử]'),
  ('Ma trận quyền','[Nhóm/danh tính; đọc/ghi/xuất/xóa; phê duyệt; ngày rà soát]'),('Mã hóa và quản lý khóa','[Khi truyền/lưu; KMS/HSM nếu dùng; chủ khóa; khôi phục/thu hồi]'),
  ('Lưu giữ và hủy','[Thời hạn; căn cứ; phương pháp; bản sao; biên bản xác minh]'),('Theo dõi và phục hồi','[Audit/DLP/toàn vẹn; job backup; RPO/RTO; phép thử]'),
  ('Bằng chứng','[Danh mục; cấu hình che bí mật; log; ticket quyền; EV]')]),
 ('PL06','Dự phòng sao lưu và khôi phục thảm họa',[
  ('Dịch vụ và mục tiêu được duyệt','[Mã PL01/BIA; RPO; RTO; MTD; người/ngày phê duyệt]'),('Phụ thuộc và thứ tự khôi phục','[Mạng/thời gian/DNS/IAM/KMS/storage/DB/queue/API/workload]'),
  ('HA và quorum','[Cụm tại từng IDC; số thành viên; failure domain; ngưỡng chịu lỗi]'),('Sao lưu','[Dữ liệu/cấu hình/khóa; job; lịch; retention; mã hóa; kho độc lập]'),
  ('Đồng bộ DC–DR','[Cơ chế; chiều; lịch; độ trễ đo; băng thông; kiểm tra nhất quán]'),('Năng lực DR và rủi ro chung','[Tải đáp ứng; nguồn điện; ISP/tuyến cáp; rủi ro đồng thời; biện pháp]'),
  ('Kích hoạt và chuyển đổi','[Ngưỡng; người quyết định; fencing; runbook; truyền thông]'),('Chuyển về và khởi động nguội','[Điều kiện; đồng bộ ngược; rollback; bí mật khẩn cấp; phép thử]'),
  ('Kết quả xác nhận','[Ngày; kịch bản; RPO/RTO đo được; đối soát; tồn tại; EV]')]),
 ('PL07','Cấu hình và phiên bản nền tảng',[
  ('Thành phần và phạm vi','[Dịch vụ; site; asset/cụm; môi trường]'),('Nguồn và bản dựng','[CẦN ĐIỀN nguồn phần mềm; phiên bản XOR; commit/tag; digest]'),
  ('Khác biệt và thành phần phụ thuộc','[Bản vá/tùy biến; SBOM; thư viện; giấy phép; hỗ trợ]'),('Baseline cấu hình','[Mã/phiên bản; hardening; IAM; log; giao thức; người duyệt]'),
  ('Chuỗi phát hành','[Kho mã; pipeline; runner; registry; người review và duyệt]'),('Kiểm tra trước phát hành','[SAST/SCA/secret scan/DAST/pentest; phạm vi; kết quả; retest]'),
  ('Triển khai và quay lui','[Ticket; ngày; cửa sổ; bản trước/sau; phương án rollback]'),('Theo dõi vận hành','[Drift; CVE; vá; EOL; tương thích và hiệu năng]'),
  ('Bằng chứng và phê duyệt','[Release manifest; log triển khai; biên bản; người ký; EV]')])
]
for code,title,rows in annexes:
 d.page('Phụ lục '+code+' '+title)
 d.p('Mã bản ghi: [CẦN ĐIỀN]. Ngày chốt hiện trạng: [CẦN ĐIỀN]. ')
 d.table(['Trường thông tin','Thông tin hiện trạng'],rows,[1.25,2.75],11.5)
 d.p('Người lập: [CẦN ĐIỀN]. Người kiểm tra: [CẦN ĐIỀN]. Văn bản xác nhận/phê duyệt: [CẦN ĐIỀN]. Tài liệu kèm theo: [mã, tên, phiên bản, số trang, nơi lưu].')
d.save(ADMIN/'01_Thuyet_minh_ho_so_de_xuat_cap_do_4.docx')

# PHỤ LỤC CHI TIẾT
d=Doc('Phụ lục đối chiếu yêu cầu và phương án bảo đảm an ninh mạng')
d.cover('Kèm theo hồ sơ đề xuất cấp độ 4 hệ thống XOR SECURE CLOUD')
d.p('Phụ lục đối chiếu 119 nội dung kiểm tra thuộc Điều 6.1–6.18 và Phụ lục A.4 TCVN 14423:2026 với phương án bảo đảm an ninh mạng của hệ thống XOR SECURE CLOUD.')
d.p('Thời điểm chốt hiện trạng: [CẦN ĐIỀN]. Hồ sơ kiểm tra và bằng chứng kèm theo: [CẦN ĐIỀN danh mục, phiên bản, người lập và người xác nhận].')
d.table(['Thông tin chung','Nội dung'],[('Phạm vi','Cả DC Cầu Giấy và DR Nam Thăng Long cùng toàn bộ dịch vụ quản trị chung.'),('Người lập','[CẦN ĐIỀN]'),('Người kiểm tra','[CẦN ĐIỀN]'),('Bản dựng và ngày chốt','[CẦN ĐIỀN]'),('Quy tắc kết luận','Đạt / Không đạt / Chưa đủ bằng chứng / Không áp dụng có căn cứ và phê duyệt.')],[1,3])
for g,title in GROUPS.items():
 items=[x for x in CONTROLS if x['group']==title]
 d.page(g.replace('.',' ')+' '+title)
 d.p('Đơn vị chủ trì: '+items[0]['owner']+'. Phạm vi hiện trạng: [CẦN XÁC NHẬN cả hai IDC và dịch vụ áp dụng].')
 rows=[]
 for x in items:
  rows.append((x['id']+'\n'+x['clause']+'\ntr.'+x['page'],x['title']+'\n'+x['requirement'],x['plan']+'\nHiện trạng: [CẦN ĐIỀN].',x['evidence']+'\nMã BC: [CẦN ĐIỀN].\nKết quả: Chưa đủ bằng chứng.'))
 d.table(['Mã và nguồn','Nội dung đối chiếu','Phương án tại XOR và hiện trạng','Bằng chứng và kết quả'],rows,[.8,2.1,2.2,1.8],10.5 if g=='6.12' else 11)
d.page('Xác nhận kết quả đối chiếu')
d.p('Tổng số điểm áp dụng: [  ]; không áp dụng có căn cứ: [  ]; đạt: [  ]; không đạt: [  ]; chưa đủ bằng chứng: [  ]. Ngày chốt và phạm vi: [CẦN ĐIỀN].')
d.p('Tồn tại chưa xử lý: [CẦN ĐIỀN mã yêu cầu, ảnh hưởng, hành động, người chịu trách nhiệm, hạn và biện pháp tạm thời].')
d.p('Kết luận của người kiểm tra: [CẦN ĐIỀN kết luận trong phạm vi và giới hạn kiểm tra thực tế].');d.sign(role='NGƯỜI CÓ THẨM QUYỀN XÁC NHẬN',signer='[CẦN ĐIỀN NGƯỜI KIỂM TRA ĐƯỢC PHÂN CÔNG]')
d.save(ADMIN/'02_Phu_luc_doi_chieu_yeu_cau_TCVN.docx')

# QUY CHẾ VÀ CÁC QUY TRÌNH
POLICIES=[
 ('QT01','Quản lý rủi ro và ngoại lệ','GRC; chủ quản phê duyệt','6.1; 6.4; 6.15','Ít nhất 6 tháng nhận diện/đánh giá hiệu lực; quy trình hằng năm và khi thay đổi.',[
 'Chủ tài sản cung cấp danh mục, dữ liệu, phụ thuộc và hậu quả gián đoạn; GRC nhận diện mối đe dọa, điểm yếu, nhà cung cấp và sự kiện có thể xảy ra.',
 'Nhóm đánh giá xác định khả năng, tác động và căn cứ; BIA xem xét cả khách hàng, lợi ích công cộng, trật tự/an toàn xã hội hoặc an ninh quốc gia nếu liên quan.',
 'Chọn tránh, giảm, chuyển giao hoặc chấp nhận theo thẩm quyền; chỉ định người, hạn, kinh phí và chỉ tiêu kiểm tra; quản lý biện pháp tương đương bằng hồ sơ riêng.',
 'Đầu mối kỹ thuật triển khai qua quản lý thay đổi; người kiểm tra xác nhận hiệu lực; GRC đánh giá tồn dư và trình chủ quản.',
 'Theo dõi thay đổi, thông báo bên liên quan và đóng có chứng cứ. Ngoại lệ có lý do, thời hạn, biện pháp bù và người duyệt; không thay thế yêu cầu bắt buộc.'
 ],'Sổ rủi ro, BIA, phiếu ngoại lệ, kế hoạch xử lý, biên bản xác nhận và quyết định chấp nhận.'),
 ('QT02','Quản lý tài sản dữ liệu và vòng đời','Hạ tầng, Platform, chủ dữ liệu','6.2–6.4','Kiểm kê hằng năm; phát hiện phần cứng/DHCP hằng quý; phần mềm trái phép 6 tháng; rà quyền dữ liệu hằng quý.',[
 'Tiếp nhận tài sản có chủ sở hữu, ID, site/rack/vùng, IP/MAC/serial, phiên bản, EOL, mục đích và bằng chứng kiểm tra trước sử dụng.',
 'Cập nhật CMDB, danh mục phần mềm và dữ liệu; phân loại độ nhạy cảm, quyền, thời hạn lưu và hủy; phần mềm phải thuộc danh sách được phép.',
 'Đối soát discovery, DHCP/IPAM, inventory nền tảng và danh mục; xử lý tài sản lạ bằng từ chối/cách ly/tiếp nhận có phê duyệt.',
 'Khi thay đổi hoặc thu hồi, cập nhật quan hệ tài sản/dữ liệu và quyền; kiểm tra snapshot, backup, image và dữ liệu còn lại theo thời hạn áp dụng.',
 'Thanh lý bằng phương pháp sanitize phù hợp media hoặc hủy vật lý; lập biên bản, xác minh và đóng vòng đời trong CMDB.'
 ],'CMDB, catalog phần mềm/SBOM, danh mục dữ liệu, ACL, phiếu tiếp nhận và biên bản hủy.'),
 ('QT03','Quản lý tài khoản và truy cập đặc quyền','IAM; chủ tài sản duyệt; bảo mật kiểm tra','6.6; 6.12.2.5–6.12.2.6','Rà tài khoản ít nhất hằng quý; rà quy trình hằng năm; thu hồi ngay theo thay đổi nhân sự.',[
 'Người yêu cầu lập phiếu nêu hệ thống, vai trò, mục đích, thời hạn và người quản lý; chủ tài sản xét quyền tối thiểu, tách nhiệm vụ.',
 'IAM tạo tài khoản cá nhân/nhóm kỹ thuật có người sở hữu; áp MFA đối với đặc quyền, bên ngoài và đối tác; tách tài khoản vận hành với tài khoản tenant.',
 'Truy cập từ PAW được bảo vệ qua gateway và PAM; chỉ đích/lệnh/quyền/thời gian được duyệt; ghi phiên và log; không cấp route rộng hoặc Internet cho tài nguyên quản trị.',
 'Thay đổi/thu hồi qua ticket, đối soát tất cả hệ thống liên quan; vô hiệu tài khoản nhàn rỗi sau 45 ngày theo điều khoản; quản lý ngoại lệ và tài khoản mặc định.',
 'Khẩn cấp dùng danh tính được kiểm soát, cảnh báo, giới hạn thời gian và hậu kiểm; người có quyền không tự phê duyệt và xóa dấu vết phiên của mình.'
 ],'Phiếu cấp/đổi/thu hồi, ma trận quyền, log MFA/PAM, rà soát quý và biên bản khẩn cấp.'),
 ('QT04','Cấu hình thay đổi và bản vá','Platform và Network; CAB hoặc người được giao duyệt','6.5; 6.7; 6.12','Scan toàn hệ thống 6 tháng, tài sản quan trọng quý; máy người dùng kiểm tra vá tháng; baseline/quy trình theo điều khoản.',[
 'Duy trì baseline theo loại tài sản, bản dựng và yêu cầu an ninh; ghi nhận port/service, xác thực, thời gian, log, mã hóa, session và cấu hình không được phép.',
 'Tiếp nhận CVE/thay đổi; đánh giá phạm vi, mức rủi ro và phụ thuộc; xác định kiểm thử, phương án rollback, cửa sổ và người phối hợp.',
 'Thử trên môi trường phù hợp, kiểm tra tương thích/hiệu năng; chỉ triển khai sau phê duyệt; nhập gói đã kiểm tra vào kho nội bộ thay vì mở Internet mạng quản trị.',
 'Thực hiện theo kế hoạch, lưu cấu hình trước/sau và log; kiểm tra chức năng/an ninh, rollback khi vượt điều kiện dừng đã chốt.',
 'Cập nhật CMDB/baseline/sơ đồ, scan và retest; ngoại lệ có hạn và kiểm soát bù; theo dõi drift và tài sản hết hỗ trợ.'
 ],'Yêu cầu thay đổi, đánh giá tác động, phê duyệt, log triển khai, rollback, scan và retest.'),
 ('QT05','Bảo vệ dữ liệu khóa và chữ ký số','Chủ dữ liệu; quản trị khóa; bảo mật','6.4; 6.11','Theo vòng đời dữ liệu/khóa; rà quyền và log dữ liệu ít nhất hằng quý.',[
 'Lập danh mục dữ liệu, chủ sở hữu, phân loại, luồng, quyền truy cập và nghĩa vụ lưu/hủy; xác định dữ liệu quan trọng và giao dịch cần chữ ký số.',
 'Áp mã hóa lưu và truyền hoặc phương án tương đương có thuyết minh; luồng quan trọng kiểm tra hai lớp mã hóa đồng thời theo yêu cầu; dùng thuật toán mạnh.',
 'Quản lý sinh, phân phối, dùng, luân chuyển, thu hồi, sao lưu và khôi phục khóa; người quản trị khóa và hạ tầng tách nhiệm vụ; không lưu bí mật trong code/hồ sơ.',
 'Giám sát toàn vẹn, DLP và truy cập dữ liệu; ghi việc xuất/sao chép; xác minh chữ ký, chứng thư và quyền khi trao đổi dữ liệu quan trọng.',
 'Hủy dữ liệu/khóa theo phân loại và thời hạn; khi lộ khóa, kích hoạt ứng phó, thu hồi, thay thế và kiểm tra ảnh hưởng bản sao.'
 ],'Danh mục dữ liệu, DFD, ma trận mật mã, log KMS, chứng thư, thử restore khóa, log DLP và biên bản hủy.'),
 ('QT06','Nhật ký giám sát và trực an ninh','SOC và SecOps','6.8; 6.10; 6.13','Trực 24/7; log tối thiểu 6 tháng; rà log tháng; điều chỉnh cảnh báo quý.',[
 'Lập ma trận tài sản/nguồn log, loại sự kiện và trường dữ liệu; đồng bộ giờ; cấu hình agent/collector có xác thực và giám sát mất nguồn.',
 'Thu log vào vùng Z08 khác vùng sinh log; giới hạn quyền, chống sửa/xóa; tính dung lượng theo số đo, cảnh báo đầy kho và xác minh retention.',
 'SOC nhận cảnh báo, phân loại, xác minh, mở ticket và chuyển chỉ huy theo mức sự cố; ca sau nhận đầy đủ công việc đang mở.',
 'Kết hợp EDR, network flow, WAF/IPS, IAM/PAM và dữ liệu ứng dụng; chỉ cô lập tài sản theo quyền hạn/playbook và lưu chứng cứ.',
 'Rà soát nguồn, quy tắc và ngưỡng; thử use case; lập báo cáo hoạt động, sự kiện thiếu nguồn và hành động khắc phục.'
 ],'Lịch trực, log bàn giao, ma trận nguồn, rule, ticket, báo cáo tháng/quý và phép thử.'),
 ('QT07','Sao lưu và khôi phục dữ liệu','Backup; chủ dịch vụ duyệt; người kiểm tra xác nhận','6.11','Tần suất theo dữ liệu/BIA; thử khôi phục ít nhất 6 tháng; quy trình hằng năm/khi đổi.',[
 'Chốt danh mục dữ liệu/config/image/DB và khóa phụ thuộc; xác định lịch, RPO/RTO, retention, phương pháp nhất quán và kho độc lập.',
 'Tự động chạy job, mã hóa dữ liệu nhạy cảm, quản lý phiên bản, quyền và giám sát lỗi; bảo vệ theo 3-2-1 hoặc phương án tương đương đã thuyết minh.',
 'Theo dõi lần thành công, khối lượng, checksum, trạng thái bản sao ngoài miền lỗi chung; điều tra lỗi và chạy bù có kiểm soát.',
 'Phục hồi vào vùng cô lập với bản sao được chọn; kiểm tra khóa, toàn vẹn và nghiệp vụ; đo dữ liệu mất/thời gian phục hồi so với mục tiêu.',
 'Lập biên bản, khắc phục và retest; điều chỉnh lịch/khả năng lưu trữ khi dữ liệu hoặc mục tiêu thay đổi.'
 ],'Catalog job, log, chứng cứ độc lập/bất biến nếu dùng, checksum và biên bản restore.'),
 ('QT08','Chuyển đổi DR và khởi động lại nền tảng','Chỉ huy DR; Platform; chủ quản quyết định','6.11; 6.12; BIA và thiết kế XOR','Theo kế hoạch được phê duyệt và khi thay đổi lớn; mục tiêu do BIA quyết định.',[
 'Xác nhận sự cố, phạm vi, trạng thái dữ liệu/replica và điều kiện kích hoạt; chỉ huy xin quyết định theo phân quyền và thông báo đầu mối liên quan.',
 'Cô lập nguồn ghi ở DC bằng cơ chế fencing được thiết kế; xác định bản dữ liệu hợp lệ; ghi thời điểm giao dịch cuối cùng có thể khôi phục.',
 'Phục hồi mạng, thời gian, DNS, IAM/khóa, lưu trữ, DB/queue, API và workload theo đồ thị phụ thuộc; dùng bí mật khẩn cấp có kiểm soát khi cần.',
 'Kiểm tra nghiệp vụ và năng lực DR, ghi RPO/RTO từng dịch vụ; mở lưu lượng theo phê duyệt, theo dõi lỗi và tính nhất quán.',
 'Chuyển về sau sửa lỗi, đồng bộ và đối soát; có điểm quyết định/rollback; diễn tập cả failback và khởi động nguội, lưu toàn bộ timeline.'
 ],'BIA, runbook, sơ đồ phụ thuộc, quyết định kích hoạt, fencing log, số đo và biên bản chuyển về.'),
 ('QT09','Ứng phó sự cố và bảo toàn chứng cứ','CSIRT; SOC; pháp chế; chủ quản','6.17','Diễn tập ít nhất hằng năm; rà quy trình/ngưỡng/danh bạ hằng năm hoặc khi đổi.',[
 'SOC tiếp nhận, mở mã sự cố, phân loại và báo đầu mối chính/dự phòng; ghi người, thời gian, tài sản, ảnh hưởng và dấu hiệu ban đầu.',
 'Chỉ huy phân công kiểm soát sự cố; cô lập theo playbook, tránh phá chứng cứ; lưu bản gốc, hash, chuỗi bàn giao và quyền truy cập.',
 'Điều tra phạm vi, nguyên nhân, dữ liệu ảnh hưởng; phối hợp IDC/ISP/đối tác và cơ quan có thẩm quyền theo nghĩa vụ áp dụng; pháp chế xác nhận thông báo bên ngoài.',
 'Loại bỏ nguyên nhân, vá/cấu hình lại, khôi phục bản sạch, đổi khóa khi cần và kiểm tra trước mở dịch vụ.',
 'Hậu kiểm timeline, tác động, quyết định và bài học; giao hành động, hạn, retest; rà lại ngưỡng và tổ chức diễn tập.'
 ],'Hồ sơ sự cố, log gốc, chuỗi chứng cứ, thông báo, quyết định, biên bản khôi phục và hậu kiểm.'),
 ('QT10','Phát triển đóng gói và phát hành XOR STACK','DevSecOps; reviewer; người duyệt production','6.16','Mỗi phát hành/thay đổi; danh mục bên thứ ba quý; đào tạo và quy trình hằng năm.',[
 'Yêu cầu và threat model xác định tài sản, ranh giới tenant, dữ liệu, API, session và yêu cầu an ninh cần kiểm thử.',
 'Phát triển trong môi trường riêng; review code; không lưu bí mật; quản lý dependency, SBOM, bản quyền và nguồn upstream.',
 'Chạy SAST/SCA/secret scan và kiểm thử chức năng/an ninh; xử lý lỗ hổng và pentest trước production theo phạm vi; ghi nhận ngoại lệ đúng thẩm quyền.',
 'Đóng gói bản dựng có commit, digest và chữ ký theo thiết kế; lưu provenance; runner có quyền hẹp; phê duyệt release và kế hoạch rollback.',
 'Triển khai, xác minh bản đang chạy, theo dõi lỗi và rollback khi cần; tiếp nhận báo cáo bên ngoài, RCA và bổ sung kiểm thử chống tái diễn.'
 ],'Yêu cầu, threat model, review, SBOM, scan/pentest, release manifest, log deploy và rollback.'),
 ('QT11','Quản lý nhân sự nhà cung cấp và vật lý','HR; mua sắm; IDC; chủ quản','6.14; 6.15; A.4','Đào tạo và rà danh mục/quy định hằng năm; theo sự kiện nhân sự/hợp đồng; camera ít nhất 6 tháng.',[
 'Chủ quản giao riêng vai trò vận hành, quản trị và an ninh; HR xác minh năng lực, NDA, đào tạo và quyền được cấp.',
 'Mua sắm phân loại nhà cung cấp, kiểm tra điều kiện tương ứng dịch vụ, phạm vi trách nhiệm và yêu cầu bảo mật trong hợp đồng.',
 'IDC và XOR quản lý người ra vào, khách, rack, điện/UPS, PCCC, nước, môi trường và điện từ; lưu camera và log theo yêu cầu; cung cấp bằng chứng từng site.',
 'Theo dõi SLA, truy cập nhà cung cấp và bảo trì; giám sát mọi phiên đặc quyền; xử lý sự cố theo đầu mối đã chốt.',
 'Khi nhân sự hoặc hợp đồng kết thúc, thu hồi thẻ/tài khoản/tài sản, bàn giao dữ liệu và log, xác nhận xóa theo nghĩa vụ; kiểm tra không còn truy cập.'
 ],'Quyết định phân công, hồ sơ đào tạo/NDA, hợp đồng/SLA, bằng chứng IDC, phiếu thu hồi.'),
 ('QT12','Kiểm tra nghiệm thu và quản lý hồ sơ','GRC; đơn vị kiểm tra; chủ quản','6.12.2.7; 6.18','Pentest ngoài và trong mỗi loại ít nhất hằng năm; kiểm tra trước vận hành/phát hành theo phạm vi.',[
 'Lập phạm vi, mục tiêu, tiêu chí, ROE, thời gian, dữ liệu thử, điều kiện dừng và đầu mối; người có thẩm quyền phê duyệt trước thử.',
 'Tổ chức kiểm tra có tư vấn/giám sát độc lập theo yêu cầu; lưu bước, cấu hình, bản dựng, log và giới hạn phương pháp.',
 'Lập phát hiện theo tài sản, tác động, bằng chứng; gắn người, hạn và hành động; không kết luận đạt ở nơi chưa đủ bằng chứng.',
 'Retest sau khắc phục, đánh giá hiệu lực kiểm soát và lập biên bản nghiệm thu có xác nhận chuyên trách, trình chủ quản phê duyệt.',
 'Đóng phiên bản hồ sơ, danh mục tài liệu/hash và quyền nhận; theo dõi phản hồi thẩm định; cập nhật khi phạm vi, dữ liệu hoặc kiến trúc thay đổi.'
 ],'ROE, báo cáo kiểm tra, findings, retest, nghiệm thu, danh mục hồ sơ và lịch sử thay đổi.')]
d=Doc('Quy chế và quy trình bảo đảm an ninh mạng XOR SECURE CLOUD')
d.cover('Kèm theo Quyết định số [CẦN ĐIỀN] ngày [CẦN ĐIỀN]')
d.p('Ban hành kèm theo Quyết định số [CẦN ĐIỀN] ngày [CẦN ĐIỀN] của '+ORG+'.')
d.h('Điều 1 Phạm vi và đối tượng áp dụng')
d.p('Quy chế áp dụng cho toàn bộ nền tảng XOR STACK CLOUD và dịch vụ quản trị chung tại IDC Cầu Giấy, IDC Nam Thăng Long; nhân sự XOR, đơn vị vận hành, đơn vị an ninh và bên thứ ba có quyền tác động đến hệ thống trong phạm vi hợp đồng.')
d.h('Điều 2 Nguyên tắc bảo vệ')
d.p('Bảo vệ theo cấp độ được phê duyệt; quản lý theo rủi ro; quyền tối thiểu; tách nhiệm vụ; phân vùng và chặn mặc định; bảo vệ dữ liệu theo phân loại; kiểm chứng bằng bằng chứng và phép thử. Các đơn vị được giao nhiệm vụ chịu trách nhiệm triển khai và duy trì hiệu lực của biện pháp bảo vệ.')
d.h('Điều 3 Trách nhiệm')
d.p('Chủ quản phê duyệt phạm vi, nguồn lực, rủi ro và quyết định theo thẩm quyền. Đơn vị vận hành triển khai, cập nhật hiện trạng và khắc phục. Đơn vị chuyên trách về an ninh mạng hướng dẫn, kiểm tra và cho ý kiến chuyên môn. GRC quản lý hồ sơ; SOC giám sát 24/7; chủ tài sản chịu trách nhiệm dữ liệu và quyền. Người thực hiện và người xác nhận phải tuân thủ cơ chế độc lập chuyên môn đã được giao.')
d.page('Điều 4 Tổ chức thực thi và lưu hồ sơ')
d.p('Mỗi quy trình QT01–QT12 phải được gắn đầu mối chính, dự phòng và danh sách hệ thống áp dụng. Kết quả thực hiện được lưu bằng mã hồ sơ/ticket, ngày, phiên bản, người lập và xác nhận. Hồ sơ có dữ liệu nhạy cảm được cấp quyền tối thiểu; bản gửi bên ngoài chỉ chứa thông tin cần thiết và được chủ quản cho phép.')
d.h('Điều 5 Xử lý trường hợp khẩn cấp và ngoại lệ')
d.p('Trường hợp khẩn cấp thực hiện theo phân quyền và điều kiện dừng đã quy định, có ghi log và hậu kiểm. Ngoại lệ phải nêu yêu cầu, lý do, rủi ro, biện pháp bù, thời hạn và người phê duyệt; được rà soát khi hết hạn hoặc thay đổi. Không dùng ngoại lệ nội bộ để thay yêu cầu bắt buộc khi quy định không cho phép.')
d.h('Điều 6 Kiểm tra và cập nhật')
d.p('Đơn vị chuyên trách theo dõi lịch kiểm tra, tình trạng bằng chứng và hành động khắc phục; báo cáo chủ quản các mục chưa đáp ứng. Quy chế và quy trình được rà soát theo tần suất tương ứng trong TCVN, khi thay đổi tổ chức/kiến trúc hoặc sau sự cố. Các mốc phản ứng, thời hạn vá và mục tiêu phục hồi phải được chốt trong phụ lục vận hành có người duyệt.')
d.table(['Tham số phải phê duyệt','Giá trị hoặc văn bản'],[
 ('Chỉ huy sự cố/DR và người dự phòng','[CẦN ĐIỀN họ tên, vai trò, liên hệ và quyền quyết định]'),
 ('Thời gian phản ứng và báo cáo sự cố','[CẦN ĐIỀN theo loại sự cố và nghĩa vụ pháp lý/hợp đồng]'),
 ('Thời hạn xử lý lỗ hổng theo rủi ro','[CẦN ĐIỀN mức, thời hạn, ngoại lệ và quyền phê duyệt]'),
 ('RPO RTO lịch backup và bất biến','[CẦN ĐIỀN theo BIA từng dịch vụ]'),
 ('Lưu giữ hồ sơ và bằng chứng','[CẦN ĐIỀN theo loại; log ANM/camera bảo đảm tối thiểu 6 tháng theo TCVN]')],[1.5,2.5])
for q,title,owner,ref,freq,steps,records in POLICIES:
 d.page(q+' '+title);d.p('Trách nhiệm: '+owner+'.');d.p('Căn cứ: TCVN '+ref+'. Tần suất: '+freq)
 for i,x in enumerate(steps):
  p=d.p('Bước '+str(i+1)+'. '+x);p.paragraph_format.keep_together=True
 d.h('Hồ sơ lưu và điều kiện hoàn thành');d.p(records)
 d.p('Hoàn thành khi các bước trong phạm vi đã thực hiện, bằng chứng có người kiểm tra và tồn tại đã được xử lý theo thẩm quyền. Đầu mối chính: [CẦN ĐIỀN]; dự phòng: [CẦN ĐIỀN]; ngày bắt đầu áp dụng: [CẦN ĐIỀN].')
d.save(ADMIN/'03_Quy_che_va_quy_trinh_bao_dam_ANM.docx')

# VĂN BẢN TRÌNH KÝ THEO CẤU TRÚC MẪU CHÍNH THỨC
def legal_intro(d):
 d.p('Căn cứ Luật An ninh mạng số 116/2025/QH15;')
 d.p('Căn cứ '+LEGAL+';')
 d.p('Căn cứ [CẦN ĐIỀN văn bản về chức năng, nhiệm vụ, thẩm quyền của cơ quan, tổ chức].')
def common_info(d):
 d.h('Phần 1 Thông tin chung')
 for s in ['1. Tên hệ thống thông tin: XOR SECURE CLOUD.','2. Đơn vị vận hành hệ thống thông tin: '+OPERATOR+'.','3. Địa chỉ: '+ADDRESS+'.','4. Cấp độ hệ thống thông tin đề xuất: Cấp độ 4.'] :d.p(s)
 d.p('Phạm vi: toàn bộ XOR STACK CLOUD và dịch vụ quản trị chung; DC tại Trung tâm Hạ tầng IDC Cầu Giấy, DR tại Trung tâm Hạ tầng IDC Nam Thăng Long.')
def attachments(d,opinion=False,appraisal=False):
 d.h('Phần 2 Hồ sơ kèm theo')
 a=['Tài liệu mô tả, thuyết minh tổng quan hệ thống thông tin.','Thiết kế thi công đã được phê duyệt hoặc tài liệu có giá trị tương đương, kèm phụ lục hiện trạng.','Thuyết minh đề xuất cấp độ có căn cứ theo Điều 14, kèm báo cáo đánh giá rủi ro chi tiết.','Thuyết minh phương án bảo đảm an ninh mạng, kèm phụ lục đối chiếu yêu cầu và quy trình.']
 if opinion:a.append('Ý kiến chuyên môn của đơn vị chuyên trách về an ninh mạng của chủ quản: số [  ], ngày [  ].')
 if appraisal:a.append('Ý kiến thẩm định của cơ quan chủ trì thẩm định và tài liệu phối hợp liên quan: số [  ], ngày [  ]; giải trình/khắc phục kèm theo nếu có.')
 for i,x in enumerate(a):d.p(f'{i+1}. '+x)
 d.p('Danh mục tệp, phiên bản và số trang: [CẦN ĐIỀN danh mục đã đóng phiên bản].')

d=Doc('Công văn xin ý kiến chuyên môn về hồ sơ đề xuất cấp độ')
d.letterhead(OPERATOR,'[KÝ HIỆU]','V/v xin ý kiến chuyên môn về hồ sơ đề xuất cấp độ 4 hệ thống XOR SECURE CLOUD','Dự thảo theo Mẫu số 03')
d.p('Kính gửi: '+SPECIALIST+' — đơn vị chuyên trách về an ninh mạng của '+ORG+'.',True)
legal_intro(d);d.p(OPERATOR+' đề nghị Quý đơn vị cho ý kiến chuyên môn về hồ sơ đề xuất cấp độ hệ thống thông tin với các nội dung sau:')
common_info(d);attachments(d)
d.p('Đề nghị Quý đơn vị xem xét sự phù hợp của đề xuất cấp độ và phương án bảo đảm an ninh mạng hệ thống XOR SECURE CLOUD; chỉ rõ nội dung cần điều chỉnh, bổ sung để đơn vị vận hành hoàn thiện trước khi trình chủ quản gửi thẩm định.')
d.p('Đầu mối phối hợp: '+AUTHOR_NAME+' — '+AUTHOR_ROLE+'; [CẦN ĐIỀN điện thoại, email].');d.sign(role=OPS_ROLE,signer=OPS_SIGNER);d.save(ADMIN/'04_Cong_van_xin_y_kien_chuyen_mon.docx')

d=Doc('Công văn đề nghị thẩm định hồ sơ đề xuất cấp độ')
d.letterhead(ORG,'[KÝ HIỆU]','V/v đề nghị thẩm định hồ sơ đề xuất cấp độ 4 hệ thống XOR SECURE CLOUD','Dự thảo theo Mẫu số 02')
d.p('Kính gửi: Bộ Công an [CẦN ĐIỀN đơn vị tiếp nhận].',True)
legal_intro(d);d.p(ORG+' đề nghị thẩm định hồ sơ đề xuất cấp độ đối với hệ thống XOR SECURE CLOUD với các nội dung sau:')
common_info(d);attachments(d,True)
d.p('Đề nghị Quý cơ quan xem xét, cho ý kiến thẩm định hồ sơ đề xuất cấp độ đối với hệ thống XOR SECURE CLOUD để chủ quản tổ chức phê duyệt và thực hiện biện pháp bảo vệ theo quy định.')
d.p('Chủ quản chịu trách nhiệm về thông tin, tài liệu gửi kèm và phối hợp cung cấp, giải trình theo yêu cầu. Đầu mối: [CẦN ĐIỀN].')
d.sign();d.save(ADMIN/'05_Cong_van_de_nghi_tham_dinh.docx')

d=Doc('Tờ trình phê duyệt hồ sơ đề xuất cấp độ')
d.letterhead(OPERATOR,'TTr-[KÝ HIỆU]',model='Dự thảo theo Mẫu số 05')
d.p('TỜ TRÌNH',True,align=WD_ALIGN_PARAGRAPH.CENTER);d.p('Về việc phê duyệt đề xuất cấp độ 4 hệ thống XOR SECURE CLOUD',True,align=WD_ALIGN_PARAGRAPH.CENTER)
d.p('Kính gửi: '+ORG+'.',True);legal_intro(d)
d.p('Căn cứ ý kiến thẩm định số [CẦN ĐIỀN] ngày [CẦN ĐIỀN] của [CƠ QUAN CHỦ TRÌ THẨM ĐỊNH]; căn cứ ý kiến chuyên môn số [  ] ngày [  ] và hồ sơ giải trình/hoàn thiện [  ];')
d.p(OPERATOR+' trình chủ quản phê duyệt hồ sơ đề xuất cấp độ với các nội dung sau:');common_info(d);attachments(d,True,True)
d.h('Phần 3 Nội dung đề nghị')
d.p('Đề nghị chủ quản xem xét, quyết định phê duyệt cấp độ 4 đối với hệ thống XOR SECURE CLOUD trong phạm vi hồ sơ kèm theo; giao đơn vị vận hành thực hiện phương án bảo đảm an ninh mạng và giao đơn vị chuyên trách kiểm tra, giám sát theo quy định.')
d.p('Tình trạng tiếp thu ý kiến thẩm định và các điều kiện còn phải thực hiện: [CẦN ĐIỀN nội dung tiếp thu, tình trạng thực hiện và mã bằng chứng].')
d.sign(role=OPS_ROLE,signer=OPS_SIGNER);d.save(ADMIN/'06_To_trinh_phe_duyet.docx')

d=Doc('Quyết định phê duyệt cấp độ hệ thống thông tin')
d.letterhead(ORG,'QĐ-[KÝ HIỆU]',model='Dự thảo theo cấu trúc Mẫu số 06')
d.p('QUYẾT ĐỊNH',True,align=WD_ALIGN_PARAGRAPH.CENTER);d.p('Về việc phê duyệt cấp độ hệ thống thông tin XOR SECURE CLOUD',True,align=WD_ALIGN_PARAGRAPH.CENTER)
d.p(ROLE,True,align=WD_ALIGN_PARAGRAPH.CENTER);legal_intro(d)
d.p('Căn cứ ý kiến thẩm định số [  ] ngày [  ] của [CƠ QUAN CÓ THẨM QUYỀN]; xét đề nghị của '+OPERATOR+' tại Tờ trình số [  ] ngày [  ];')
d.p('QUYẾT ĐỊNH',True,align=WD_ALIGN_PARAGRAPH.CENTER)
d.h('Điều 1 Phê duyệt cấp độ hệ thống thông tin')
d.p('1. Tên hệ thống thông tin: XOR SECURE CLOUD. Đơn vị vận hành: '+OPERATOR+'. Địa chỉ: '+ADDRESS+'. Phạm vi hệ thống và các thành phần được xác định tại hồ sơ mã [  ], phiên bản [  ], ngày [  ] kèm theo Quyết định này.')
d.p('2. Cấp độ hệ thống thông tin: Cấp độ 4.')
d.p('3. Phương án bảo đảm an ninh mạng trong thiết kế và trong vận hành: [CẦN ĐIỀN kết luận phù hợp với tiêu chuẩn quốc gia/quy chuẩn áp dụng, đúng nội dung ý kiến thẩm định và hồ sơ đã hoàn thiện; tên, phiên bản và phạm vi tài liệu kèm theo].')
d.p('4. Điều kiện, nội dung phải hoàn thiện hoặc yêu cầu duy trì theo ý kiến thẩm định: [CẦN ĐIỀN điều kiện và yêu cầu theo ý kiến thẩm định].')
d.h('Điều 2 Tổ chức thực hiện')
d.p('1. '+OPERATOR+' chịu trách nhiệm tổ chức thực hiện phương án bảo đảm an ninh mạng; duy trì các biện pháp quản lý, kỹ thuật, vật lý; kiểm tra, khắc phục và cập nhật hồ sơ theo quy định.')
d.p('2. '+SPECIALIST+' là đơn vị chuyên trách về an ninh mạng, kiểm tra, giám sát việc thực hiện; báo cáo chủ quản các thay đổi, rủi ro và tồn tại cần xử lý.')
d.p('3. [CÁC ĐƠN VỊ LIÊN QUAN] thực hiện trách nhiệm theo phân công và hợp đồng; phối hợp cung cấp bằng chứng, ứng phó sự cố và bảo đảm liên tục.')
d.h('Điều 3 Điều khoản thi hành')
d.p('Quyết định có hiệu lực kể từ ngày [CẦN ĐIỀN]. Thủ trưởng các đơn vị nêu tại Điều 2 và các tổ chức, cá nhân liên quan chịu trách nhiệm thi hành Quyết định này.')
d.sign(role=ROLE);d.save(ADMIN/'07_Quyet_dinh_phe_duyet_cap_do.docx')

exec((TMP/'official_forms.py').read_text(), globals())

exec((TMP/'supplement.py').read_text(), globals())

for folder in [TECH,ADMIN]:
 for name in ['01_Topo_tong_the_DC_DR','02_Topo_phan_vung_va_luong']:
  for ext in ['svg','png']:shutil.copy2(TMP/(name+'.'+ext),folder/(name+'.'+ext))
(TMP/'doc_manifest.json').write_text(json.dumps(CREATED,ensure_ascii=False,indent=2),encoding='utf-8')
assert len(CREATED)==15,len(CREATED)
print(json.dumps({'docx_count':len(CREATED),'files':CREATED},ensure_ascii=False,indent=2))
