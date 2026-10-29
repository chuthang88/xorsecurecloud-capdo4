from pathlib import Path
from copy import deepcopy
import json, shutil
from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from content import *

TECH=OUT/'Bo-1-Ky-thuat'; ADMIN=OUT/'Bo-2-Hanh-chinh'
TECH.mkdir(parents=True,exist_ok=True);ADMIN.mkdir(parents=True,exist_ok=True)
CREATED=[]
DRAFT=''

class Doc:
 def __init__(self,title,formal=True):
  self.doc=Document();self.formal=formal;self.title=title
  self.font='Times New Roman' if formal else 'Arial';self.size=13 if formal else 11
  s=self.doc.sections[0]
  if formal:s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=Cm(3);s.right_margin=Cm(2);s.top_margin=s.bottom_margin=Cm(2)
  else:s.page_width=Inches(8.5);s.page_height=Inches(11);s.left_margin=s.right_margin=Inches(.7);s.top_margin=s.bottom_margin=Inches(.65)
  self.width=s.page_width-s.left_margin-s.right_margin
  s.header_distance=s.footer_distance=Cm(.7)
  for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
   st=self.doc.styles[name];st.font.name=self.font;st.font.size=Pt(self.size);st.font.color.rgb=RGBColor(0,0,0)
   fonts=st.element.get_or_add_rPr().get_or_add_rFonts()
   for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
    if qn('w:'+attr) in fonts.attrib:del fonts.attrib[qn('w:'+attr)]
   for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),self.font)
   st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.1
   for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
  for name,size in [('Title',19 if formal else 23),('Heading 1',15),('Heading 2',13),('Heading 3',12)]:
   st=self.doc.styles[name];st.font.size=Pt(size);st.font.bold=True
  self.doc.styles['Caption'].font.size=Pt(10)
  header=s.header.paragraphs[0];header.text=DRAFT
  for r in header.runs:r.font.name=self.font;r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
  f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
  f.add_run('XOR SECURE CLOUD  •  Phiên bản 2.5  •  ').font.size=Pt(9)
  fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
  self.doc.core_properties.title=title;self.doc.core_properties.author='XOR CLOUD';self.doc.core_properties.subject='Hồ sơ lần đầu đề xuất cấp độ 4'
 def p(self,text='',bold=False,style=None,align=None):
  p=self.doc.add_paragraph(style=style);p.add_run(text).bold=True if style=='Title' else bold
  if align is not None:p.alignment=align
  return p
 def h(self,text,level=2):return self.doc.add_heading(text,level)
 def page(self,title):
  p=self.h(title,1);p.paragraph_format.page_break_before=True;return p
 def cover(self,subtitle):
  self.p('XOR SECURE CLOUD',True,align=WD_ALIGN_PARAGRAPH.CENTER)
  self.p(self.title,style='Title',align=WD_ALIGN_PARAGRAPH.CENTER)
  self.p(subtitle,align=WD_ALIGN_PARAGRAPH.CENTER)
  self.p('Nền tảng XOR STACK CLOUD',align=WD_ALIGN_PARAGRAPH.CENTER)
  self.p('DC tại Trung tâm Hạ tầng IDC Cầu Giấy\nDR tại Trung tâm Hạ tầng IDC Nam Thăng Long',align=WD_ALIGN_PARAGRAPH.CENTER)
  self.p('Phiên bản 2.5 ngày 17 tháng 10 năm 2026',align=WD_ALIGN_PARAGRAPH.CENTER)
 def table(self,headers,rows,widths=None,size=None):
  size=size or (11.5 if self.formal else 10)
  t=self.doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
  if widths is None:widths=[1]*len(headers)
  scale=self.width/sum(widths)
  for i,w in enumerate(widths):t.columns[i].width=int(w*scale)
  for i,x in enumerate(headers):t.rows[0].cells[i].text=str(x)
  for row in rows:
   for c,x in zip(t.add_row().cells,row):c.text=str(x)
  for ri,row in enumerate(t.rows):
   tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
   if ri==0:tr.append(OxmlElement('w:tblHeader'))
   for ci,c in enumerate(row.cells):
    c.width=int(widths[ci]*scale);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9ECEF' if self.formal and ri==0 else ('203E59' if ri==0 else ('F5F7F9' if ri%2 else 'FFFFFF')));tc.append(sh)
    mar=OxmlElement('w:tcMar')
    for side in ['top','left','bottom','right']:
     e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
    tc.append(mar);bs=OxmlElement('w:tcBorders')
    for side in ['top','bottom','left','right']:
     b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'BFC6CC');bs.append(b)
    tc.append(bs)
    for p in c.paragraphs:
     p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.03
     for r in p.runs:
      r.font.name=self.font;r.font.size=Pt(size);r.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 and not self.formal else '000000')
  self.p('').paragraph_format.space_after=Pt(1)
  return t
 def sign(self,recipient='Như trên;\nLưu: hồ sơ XOR SECURE CLOUD.',role=ROLE,signer=SIGNER):
  for p in self.doc.paragraphs[-2:]:p.paragraph_format.keep_with_next=True
  t=self.doc.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=int(self.width*.4);t.columns[1].width=int(self.width*.6)
  t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  t.cell(0,0).text='Nơi nhận:\n'+recipient;t.cell(0,1).text=role+'\n\n\n\n'+signer
  for i,c in enumerate(t.rows[0].cells):
   for p in c.paragraphs:
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.name=self.font;r.font.size=Pt(11)
  return t
 def letterhead(self,sender,ref,subject=None,model=None):
  if model:self.p(model,align=WD_ALIGN_PARAGRAPH.RIGHT)
  t=self.doc.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=int(self.width*.4);t.columns[1].width=int(self.width*.6)
  t.cell(0,0).text=sender+'\nSố: [CẦN ĐIỀN]/'+ref
  t.cell(0,1).text='CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n\nHà Nội, ngày [  ] tháng [  ] năm [  ]'
  for c in t.rows[0].cells:
   for p in c.paragraphs:
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(5)
    for r in p.runs:r.font.name=self.font;r.font.size=Pt(11.5);r.bold=True
  if subject:self.p(subject,True)
 def save(self,path):self.doc.save(path);CREATED.append(str(path));return path

def sources(d,cloud=False):
 d.h('Tài liệu tham chiếu')
 d.p('[S1] Nghị định 331/2026/NĐ-CP: Điều 14, 18, 20–24 và Phụ lục. Nguồn chính thức: '+URL)
 d.p('[S2] TCVN 14423:2026: Điều 6, trang 58–81; Phụ lục A.4, trang 109–111.')
 if cloud:d.p('Dự thảo TCVN Cloud năm 2026 chỉ dùng để rà soát thiết kế bổ sung. Các điểm bổ sung không tạo kết luận Hạng B/Mức 4 và không thay checklist toàn bộ dự thảo.')

from doc_style import restyle_doc_class
Doc=restyle_doc_class(Doc)

