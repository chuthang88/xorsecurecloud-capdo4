from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).parent
OUT = ROOT / 'Huong dan su dung Horizon va Xor Secure Cloud.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
sec.footer_distance = Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    style=doc.styles[name]
    style.font.name='Liberation Sans'
    fonts=style.element.get_or_add_rPr().find(qn('w:rFonts'))
    for attr in list(fonts.attrib):
        if attr.endswith('Theme') or attr.endswith('theme'):
            del fonts.attrib[attr]
    style.font.color.rgb=RGBColor(0,0,0)
    style.font.size=Pt(11)
    style.paragraph_format.space_after=Pt(7)
    style.paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
doc.styles['Caption'].font.size=Pt(9)
doc.styles['Caption'].font.italic=False
doc.styles['Caption'].paragraph_format.space_after=Pt(9)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Horizon và Xor Secure Cloud  |  '); r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.core_properties.title='Hướng dẫn sử dụng Horizon và Xor Secure Cloud'
doc.core_properties.subject='Tra cứu tài nguyên và giám sát qua JumpServer'
doc.core_properties.author='XOR Secure Cloud'

def p(text, boldlead=None):
    v=doc.add_paragraph()
    if boldlead and text.startswith(boldlead):
        v.add_run(boldlead).bold=True; v.add_run(text[len(boldlead):])
    else: v.add_run(text)
    return v

def h(text): doc.add_heading(text,2)
def page(num,title):
    v=doc.add_heading(f'{num}  {title}',1)
    v.paragraph_format.page_break_before=True
def steps(items):
    for i,t in enumerate(items,1):
        v=p(f'{i}. {t}'); v.paragraph_format.space_after=Pt(5)
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    if widths:
        for col,w in zip(t.columns,widths): col.width=Inches(w)
    for i,x in enumerate(headers): t.rows[0].cells[i].text=x
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): cells[i].text=x
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); keep=OxmlElement('w:cantSplit'); trpr.append(keep)
        if ri==0:
            rep=OxmlElement('w:tblHeader'); trpr.append(rep)
        for ci,c in enumerate(row.cells):
            if widths: c.width=Inches(widths[ci])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EEF4' if ri==0 else ('F7F9FB' if ri%2==0 else 'FFFFFF')); pr.append(sh)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                edgeel=OxmlElement('w:'+edge); edgeel.set(qn('w:val'),'single'); edgeel.set(qn('w:sz'),'4'); edgeel.set(qn('w:color'),'D9D9D9'); borders.append(edgeel)
            pr.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:w'),'85'); el.set(qn('w:type'),'dxa'); mar.append(el)
            pr.append(mar)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2); para.paragraph_format.line_spacing=1.05
                for r in para.runs: r.font.size=Pt(10); r.bold=ri==0
    p('').paragraph_format.space_after=Pt(1)
    return t

fig=0
def picture(name,caption):
    global fig
    path=ROOT/'survey'/f'{name}.png'
    assert path.exists(),path
    fig+=1
    v=doc.add_paragraph(); v.paragraph_format.space_after=Pt(3); v.paragraph_format.keep_with_next=True
    v.alignment=WD_ALIGN_PARAGRAPH.CENTER
    shape=v.add_run().add_picture(str(path),width=Inches(7.1))
    shape._inline.docPr.set('descr',caption)
    doc.add_paragraph(f'Hình {fig}. {caption}',style='Caption')

doc.add_paragraph('Hướng dẫn sử dụng Horizon\nvà Xor Secure Cloud',style='Title')
doc.add_paragraph('Tra cứu tài nguyên và giám sát qua JumpServer',style='Subtitle')
p('Phiên bản 1.0  |  Ngày khảo sát 13 tháng 10 năm 2026  |  Tài liệu nội bộ')
p('Tài liệu hướng dẫn người dùng và nhân sự vận hành tìm tài nguyên, đọc trạng thái, xem cấu hình và tra cứu cảnh báo trên hai portal đang được truy cập qua ứng dụng JumpServer. Các quy trình trong tài liệu phục vụ tra cứu và giám sát, không yêu cầu thay đổi tài nguyên đang chạy.')
p('Giao diện minh họa được mở bằng tài khoản admin trong project admin, domain Default. Tài khoản của người dùng có thể thấy ít menu hoặc ít nút hơn tùy quyền được cấp. Không dùng chung tài khoản admin chỉ để làm theo ảnh minh họa.')
h('Cách sử dụng tài liệu')
p('Bắt đầu từ phần 1 để xác định đúng portal và project. Chọn phần theo tài nguyên cần xem. Khi cần hỗ trợ, ghi lại tên, ID tài nguyên, trạng thái, thời điểm và thông báo lỗi theo mẫu ở phần 18.')
table(['Nhu cầu','Phần'],[
('Truy cập và thao tác tra cứu chung','1'),('Máy ảo trên XOR và Horizon','2 và 3'),('Volume, snapshot, mạng và Security Group','4 đến 6'),('Giám sát và nhật ký','7'),('Security Center và cảnh báo bảo mật','8 đến 12'),('Hồ sơ thiết kế và các trang TCVN','13'),('Database, Kubernetes và các chức năng khác','14'),('Horizon dành cho vận hành','15'),('Đối chiếu portal, xử lý vướng mắc và yêu cầu hỗ trợ','16 đến 18')],[5.8,1.3])
p('Các con số trong ảnh là ví dụ tại thời điểm khảo sát, không phải trạng thái hiện thời hay cam kết năng lực. Những chức năng tạo, sửa, triển khai, quét và ứng cứu cần quy trình riêng trước khi thực hiện.')

page(1,'Truy cập đúng portal và project')
steps(['Mở ứng dụng JumpServer trên Mac và sử dụng phiên truy cập đã được đơn vị vận hành cấp quyền.', 'Nếu hai phiên đã mở, chọn tab WEB-XOR-Portal hoặc WEB-Horizon ở thanh trên cùng. Danh mục tài sản nằm trong XOR Infrastructure → Web Management.', 'Kiểm tra tên portal, project và domain trước khi đọc dữ liệu. XOR đang hiển thị admin và Default; Horizon hiển thị bộ chọn project ở góc trên.', 'Chờ trang tải xong. Nếu đang hiện Loading hoặc No Data trong lúc tải, chưa kết luận tài nguyên bị mất.'])
table(['Tab JumpServer','Địa chỉ hiển thị trong phiên'],[('WEB-XOR-Portal','http://10.3.10.100:9999/'),('WEB-Horizon','http://10.3.10.100/')],[2.2,4.9])
picture('01-xor-overview','Trang Home của XOR với thông tin tài khoản và quota')
p('Địa chỉ trên dùng để nhận diện phiên nội bộ. Trong môi trường này, hai portal được truy cập bên trong JumpServer. Thanh địa chỉ của phiên web đang ở chế độ chỉ đọc.')
p('Tra cứu an toàn: mở menu, chọn liên kết tên hoặc ID tài nguyên, đổi tab thông tin, lọc danh sách và xem lịch sử. Không chọn checkbox hàng loạt rồi sử dụng nút hành động.')

page(2,'Tra cứu máy ảo trên XOR')
steps(['Chọn Compute → Instances.', 'Tìm hàng theo tên hoặc ID. Ô lọc hỗ trợ nhiều thẻ, phân cách bằng Enter; dùng đúng tiêu chí mà giao diện gợi ý.', 'Đọc Status, Fixed IP, Floating IP, Flavor và Image. Bấm liên kết ID để mở Instance Detail.', 'Trong Detail, đối chiếu Name, ID, Host, Network Info, Flavor Info, Image Info và Security Group Info.'])
picture('03-xor-instance-detail','Instance Detail và các tab thông tin của một máy ảo')
table(['Tab','Thông tin cần xem'],[('Volumes và Instance Snapshots','Ổ đĩa gắn với máy ảo và các snapshot liên quan.'),('Interfaces và Floating IPs','Cổng mạng, địa chỉ nội bộ và địa chỉ Floating IP.'),('Security Groups','Nhóm quy tắc mạng gắn với máy ảo.'),('Action Logs','Operation Name, Operation Time và Request ID để đối chiếu lịch sử.')],[2,5.1])
p('Active là trạng thái tài nguyên trên cloud, không tự chứng minh ứng dụng bên trong đang phục vụ tốt. Console, Start, Stop, Reboot, Attach Volume và các mục More là chức năng vận hành; không dùng trong quy trình tra cứu này.')

page(3,'Tra cứu máy ảo trên Horizon')
steps(['Chọn WEB-Horizon → Project → Compute → Instances.', 'Chọn tiêu chí ở ô lọc, nhập tên hoặc ID phù hợp rồi bấm Filter. Kiểm tra project trước nếu kết quả không đúng mong đợi.', 'Đọc Status, Task, Power State, Availability Zone, Flavor và IP Address.', 'Bấm tên máy ảo để xem Overview. Dùng ID để đối chiếu với XOR, thay vì chỉ dựa vào tên.'])
picture('29-horizon-instances','Danh sách Instances trên Horizon trong cùng project')
table(['Mục chi tiết','Cách đọc'],[('Overview','Thông tin nhận diện, cấu hình, IP, Security Groups, metadata và volume gắn kèm.'),('Interfaces','Tra cứu các cổng mạng của máy ảo.'),('Action Log','Tra cứu lịch sử yêu cầu thao tác. Ghi lại Request ID khi cần hỗ trợ.'),('Log và Console','Có trong giao diện. Console là phiên tương tác với máy ảo; không mở trong quy trình chỉ xem này.')],[1.6,5.5])
p('Ví dụ khảo sát: cùng máy xor-test-public29 có cùng ID trên hai portal. Horizon hiển thị Active, Task None và Power State Running. Không dùng Create Snapshot, Launch Instance, Delete Instances hoặc menu hành động để “kiểm tra thử”.')

page(4,'Tra cứu volume và bản sao dữ liệu')
steps(['Trên XOR, vào Storage → Volumes; trên Horizon, vào Project → Volumes → Volumes.', 'Tìm volume theo ID hoặc tên. Đọc Size, Status, Type, Attached To và Bootable.', 'Nếu volume đang In-use, xem máy ảo và thiết bị được ghi trong Attached To. Có thể mở liên kết máy ảo để đối chiếu.', 'Với bản sao, mở Volume Snapshots hoặc Volume Backups trên XOR; dùng Snapshots hoặc Backups trong nhóm Volumes của Horizon.'])
picture('32-horizon-volumes','Horizon hiển thị volume cùng máy ảo và thiết bị đang gắn')
table(['Trường hoặc mục','Ý nghĩa trong tra cứu'],[('In-use / Attached To','Volume đang được gắn; xác định đúng máy ảo trước khi gửi yêu cầu xử lý.'),('Bootable','Volume có thuộc tính dùng để khởi động. Không đồng nghĩa đã có bản sao khôi phục.'),('Encrypted','Horizon có cột thể hiện thuộc tính mã hóa volume.'),('Snapshot và Backup','Là các loại bản sao khác nhau trong menu. Có bản ghi không đồng nghĩa đã kiểm thử khôi phục.')],[2,5.1])
p('Quota dung lượng tổng hợp có thể khác dung lượng của một volume đang xem. Không dùng Delete, Extend, Attach/Detach, Accept Volume Transfer, Restore hoặc tạo lịch sao lưu trong bước tra cứu.')

page(5,'Tra cứu mạng và sơ đồ kết nối')
steps(['XOR: Network → Networks. Chọn Current Project Networks, Shared Networks, External Networks hoặc All Networks theo nhu cầu.', 'Bấm ID mạng để xem Detail, Subnets và Ports. Trong Subnets, đọc CIDR, Gateway IP, IP Version, Total IPs và Used IPs.', 'Horizon: Project → Network → Network Topology để xem quan hệ giữa mạng, router và máy ảo.', 'Dùng Project → Network → Networks hoặc Routers khi cần xem danh sách thay cho sơ đồ.'])
picture('31-horizon-topology','Network Topology trên Horizon với mạng ngoài và các mạng project')
p('Trong sơ đồ, mở View Instance Details khi cần thông tin máy ảo. Không chọn Open Console. Các nút Launch Instance, Create Network và Create Router tạo tài nguyên mới, không phục vụ tra cứu.')
p('External và Shared là hai thuộc tính riêng của mạng. Để xác định đường đi, đối chiếu mạng của máy ảo, subnet, gateway và router; không suy luận chỉ từ tên mạng. Trạng thái Active không chứng minh toàn bộ luồng kết nối ứng dụng đã thông suốt.')

page(6,'Tra cứu Floating IP và Security Group')
h('Đối chiếu địa chỉ IP')
steps(['XOR: Network → Floating IPs. Xem ID/Floating IP, Associated Resource và Status.', 'So sánh Associated Resource với máy ảo và Fixed IP trong Compute → Instances.', 'Horizon có mục tương ứng tại Project → Network → Floating IPs. Ghi đầy đủ IP và ID khi gửi yêu cầu hỗ trợ.'])
h('Đọc quy tắc mạng')
steps(['XOR: Network → Security Groups → bấm ID nhóm → Rules.', 'Đọc Direction, Ether Type, IP Protocol, Port Range và Remote IP Prefix hoặc Remote Group Id.', 'Đối chiếu nhóm đang gắn ở Instance Detail → Security Groups. Horizon cũng có mục Project → Network → Security Groups.'])
picture('08-xor-security-rules','Các trường dùng để đọc quy tắc Security Group trên XOR')
p('Ingress là lưu lượng vào, Egress là lưu lượng ra. IPv4 và IPv6 có quy tắc riêng. Any không phải một cổng cụ thể. Mục 0.0.0.0/0 biểu diễn toàn bộ dải IPv4; cần đọc cùng chiều và giao thức.')
p('Không dùng Allocate IP, Release, Associate/Disassociate hoặc Create Rule/Delete để kiểm tra kết nối. Nếu nghi ngờ quy tắc chưa phù hợp, gửi yêu cầu thay đổi kèm nguồn, đích, giao thức và cổng cần dùng.')

page(7,'Theo dõi quota giám sát và nhật ký')
h('Quota và mức sử dụng')
p('XOR Home hiển thị quota theo Compute, Network, Storage, Clusters và Database. Horizon có Project → Compute → Overview → Limit Summary. Đọc cả tử số đã dùng và mẫu số giới hạn, không chỉ màu cảnh báo.')
p('Horizon Usage Summary dùng khoảng From và To theo định dạng YYYY-MM-DD. Kết quả theo kỳ cần được phân biệt với danh sách tài nguyên hiện tại. Trong phiên khảo sát, Limit Summary hiển thị 5 máy ảo còn Usage Summary cho kỳ đã chọn hiển thị 7 dòng.')
picture('10-xor-monitor','Monitor Overview hiển thị thông tin theo project')
h('Giám sát và lịch sử')
p('Vào Monitor Center → Monitor Overview để xem tổng quan. Alarm Rules, Notification Channels và Alarm History là các mục liên quan. Màn hình ghi rõ đây là thông tin project, không phải thống kê hypervisor.')
p('Vào Maintenance → Audit Logs để xem thời gian, mã kết quả, người dùng và độ trễ yêu cầu. Với một máy ảo cụ thể, ưu tiên Action Logs trong Instance Detail. Khi lập yêu cầu hỗ trợ, ghi nguyên chuỗi thời gian và múi giờ nếu biết; hai portal có thể hiển thị giờ khác nhau.')
p('Nếu số đếm trên Monitor Overview khác danh sách chuyên biệt, đối chiếu project, thời điểm tải và nguồn dữ liệu trước. Không tạo Alarm Rule hoặc đổi Notification Channel chỉ để làm xuất hiện biểu đồ.')

page(8,'Định hướng sử dụng Security Center')
p('Mở Security Center ở menu trái của XOR. Bắt đầu ở Overview của dịch vụ cần xem; kiểm tra Status, Health, số connector, phạm vi dữ liệu và thời điểm báo cáo gần nhất trước khi đọc số cảnh báo.')
table(['Dịch vụ','Mục đích tra cứu','Các tab hiện có'],[
('Security Posture','Kiểm kê và phát hiện cấu hình','Overview, Assets, Scans, Findings, Policies, Remediation'),
('Web Application Firewall','Ứng dụng được bảo vệ và sự kiện WAF','Overview, Applications, Policies, Events, Certificates, Usage, Settings'),
('Vulnerability Management','Tài sản, CVE và kết quả quét','Overview, Connectors, Assets, CVEs, Scans, Exceptions, Remediation, Settings'),
('SIEM','Nguồn log, cảnh báo và sự cố','Overview, Log Sources, Alerts, Incidents, Settings'),
('Endpoint Security','Máy được giám sát và phát hiện','Overview, Connectors, Endpoints, Detections, Policies, Isolation, Deployment, Settings'),
('Network Detection','Cảm biến và phát hiện mạng','Overview, Connectors, Sensors, Flows, Detections, Rules, Topology, Settings'),
('Data Protection','Nguồn dữ liệu và kết quả phân loại','Overview, Connectors, Data Sources, Classification, Findings, Policies, Actions, Settings'),
('Automation','Playbook và lịch sử thực thi','Overview, Playbooks, Runs, Integrations, Approvals, Settings'),
('Network Access Control','Thiết bị và trạng thái truy cập mạng','Overview, Connectors, Devices, Policies, Segments, Guests, Quarantine, Settings')],[1.65,2.05,3.4])
h('Đọc trạng thái đúng ngữ cảnh')
p('healthy mô tả tình trạng được dịch vụ báo cáo; cần đọc cùng thời điểm heartbeat. not_connected hoặc offline có thể làm số cảnh báo bằng 0 vì chưa có dữ liệu. Not enabled là trạng thái chưa bật dịch vụ. Không xem các trạng thái này là kết luận hệ thống an toàn.')
p('Acknowledge, Resolve, Plan, Enable, Disable, Save, Approve, Isolation và Quarantine đều có thể cập nhật trạng thái hoặc kích hoạt hành động. Khi chỉ tra cứu, dừng ở việc đọc nội dung.')

page(9,'Đọc kết quả Security Posture')
steps(['Vào Security Center → Security Posture → Overview.', 'Đọc Security Score, Assets, Open Findings, phân bố severity và Latest scan.', 'Kiểm tra thời điểm scan và độ bao phủ kiểm tra. Màn hình nêu rõ kiểm tra không có kết quả không được coi là PASS.', 'Chọn Findings. Lọc theo severity và trạng thái; đọc tên phát hiện, mô tả bằng chứng, tài sản, Last Seen và Recommendation.'])
picture('11-xor-posture','Tổng quan Security Posture với kết quả quét đã có')
p('Trong phiên khảo sát, giao diện hiển thị scan hoàn thành, score 44, 19 tài sản và 38 phát hiện mở; đây là kết quả của lần quét được ghi trên màn hình, không phải kiểm tra mới tại thời điểm đọc tài liệu.')
p('Recommendation là đề xuất của bộ kiểm tra, cần được người phụ trách đối chiếu với kiến trúc và ngoại lệ đã phê duyệt. Không áp dụng trực tiếp lên hệ thống chỉ vì có cảnh báo.')
p('Run posture scan có thể tạo tác vụ quét và phát sinh tải dù dùng API đọc. Acknowledge, Plan, Resolve và thay đổi Policies/Remediation cập nhật dữ liệu quản lý; các bước đó không thuộc quy trình tra cứu.')

page(10,'Tra cứu lỗ hổng và bảo vệ đầu cuối')
h('Vulnerability Management')
steps(['Mở Overview và kiểm tra Health, Connectors, Assets cùng Last Connector Report.', 'Dùng Assets để định vị tài sản; dùng CVEs để tra cứu phát hiện theo tài sản và mức độ nghiêm trọng.', 'Dùng Scans để xem thông tin quét đã có. Đối chiếu phạm vi quét trước khi diễn giải số CVE. Không kích hoạt quét mới hoặc thêm connector.'])
picture('14-xor-vulnerability','Tổng quan Vulnerability Management theo dữ liệu connector')
h('Endpoint Security')
steps(['Mở Overview, kiểm tra Health, số Endpoints, Open Detections và Last Manager Report.', 'Dùng Endpoints để tra cứu máy được ghi nhận; dùng Detections để xem phát hiện.', 'Ghi nhận tài sản, mức độ, trạng thái và thời điểm; chuyển người phụ trách xác minh trước mọi hành động ứng cứu.'])
p('Các số đếm CVE và detection thể hiện dữ liệu mà portal nhận được, không tự đồng nghĩa với số máy bị xâm nhập. Màn hình mô tả việc giữ telemetry thô tại manager hoặc connector; portal tập trung vào metadata đã chuẩn hóa.')
p('Không dùng Add connector, Add manager connector, Isolation, Deployment hoặc thay đổi Policies/Settings trong bước chỉ xem.')

page(11,'Tra cứu SIEM và sự cố')
steps(['Vào Security Center → SIEM → Overview, kiểm tra Connected Sources và Last Heartbeat.', 'Mở Log Sources để xác định nguồn; mở Alerts để xem cảnh báo đã được chuẩn hóa.', 'Ở Alerts, dùng All severities, All statuses và Source type để thu hẹp kết quả.', 'Đọc Alert, Severity, Source, Status, Last Seen; chuyển sang Incidents khi cần xem bản ghi sự cố liên quan.'])
picture('16-xor-siem-alerts','Các bộ lọc và trường thông tin cảnh báo SIEM')
p('Portal ghi rõ raw log nằm tại dịch vụ Wazuh riêng. Nội dung Alert trên portal là metadata, vì vậy cần chuyển người phụ trách SIEM khi cần log gốc để xác minh.')
p('Không bấm Acknowledge hoặc Resolve khi chỉ tra cứu vì hai nút này cập nhật trạng thái cảnh báo. Nhiều cảnh báo giống nhau cần được đối chiếu nguồn và thời gian; không tự coi chúng là các sự cố độc lập.')
p('Nếu Last Heartbeat đã cũ hoặc nguồn không kết nối, số cảnh báo bằng 0 không đủ để kết luận không có sự kiện. Ghi nhận trạng thái thu thập cùng nội dung yêu cầu hỗ trợ.')

page(12,'Tra cứu các dịch vụ bảo vệ bổ sung')
table(['Dịch vụ','Cách tra cứu','Điểm cần phân biệt'],[
('WAF','Overview → Applications để xem phạm vi; Events để đọc sự kiện; Certificates để tra cứu metadata chứng thư.','Đọc gateway và Health trước số blocked/allowed. Không mở hoặc xuất khóa riêng.'),
('Network Detection','Overview → Sensors/Flows/Detections; dùng Topology để xem thông tin mạng do dịch vụ cung cấp.','Số flow và detection phụ thuộc cảm biến đã kết nối và cửa sổ thời gian.'),
('Data Protection','Overview → Data Sources/Classification/Findings.','Số lượng và nhãn phân loại không thay thế nội dung bằng chứng nằm ở connector.'),
('Network Access Control','Overview → Devices/Segments; đọc Collection Status.','Collector healthy hoặc engine action thành công không tự chứng minh switch đã thực thi chính sách.'),
('Automation','Overview → Playbooks/Runs; xem trạng thái runner và kết quả có sẵn.','Request playbook run và Approvals có thể dẫn đến thực thi, không phải thao tác xem.')],[1.55,2.7,2.85])
picture('13-xor-waf','WAF Overview và trạng thái gateway tại thời điểm khảo sát')
p('Tại thời điểm khảo sát, WAF hiển thị error/unknown và không có gateway hoạt động được đăng ký; NDR, DLP, NAC hiển thị not_connected; Automation hiển thị offline. Đây là trạng thái cần đối chiếu với cấu hình dịch vụ, không phải kết luận về toàn bộ hệ thống.')

page(13,'Xem hồ sơ thiết kế và các trang TCVN')
p('Mở X-OR SECURE → Hồ sơ thiết kế. Bộ chọn Chọn hồ sơ dùng để xác định hồ sơ đang xem. Chú ý các nhãn Chưa lưu trên server và Hồ sơ mới để không nhầm mẫu biểu với hồ sơ đã được ghi nhận.')
picture('22-xor-design','Hồ sơ thiết kế với các bước và thông tin phạm vi')
table(['Nhóm bước trong giao diện','Nội dung tra cứu'],[('1 đến 3','Thông tin hệ thống; Khảo sát và tài sản; Topology và mạng.'),('4 đến 6','Chính sách truy cập; Giải pháp và tài nguyên; Đối chiếu yêu cầu.'),('7 và 8','Kiểm thử và vận hành; Trình duyệt thiết kế.')],[2,5.1])
p('Các mục TCVN Cấp độ 1, 2 và 3 mở trang theo cấp độ lựa chọn. Trang Cấp độ 3 có các khu vực Tổng quan, VM cần bảo vệ, Kiến trúc, Topo thực tế, Kiểm soát TCVN và Triển khai.')
p('Đây là hướng dẫn sử dụng giao diện. Đề xuất số VM, điểm hiển thị hoặc phê duyệt thiết kế không phải chứng nhận đạt cấp độ. Giao diện cũng nêu số VM là phương án của hệ thống, không phải số lượng do TCVN ấn định.')
p('Khi chỉ xem, không nhập form, không chọn Tạo bản thiết kế, Lưu bản nháp, Thiết lập & kích hoạt hoặc xác nhận triển khai. Nếu cần lập hồ sơ chính thức, dùng quy trình thiết kế và rà soát căn cứ riêng.')

page(14,'Tra cứu các dịch vụ cloud khác')
table(['Nhóm chức năng XOR','Đường đi và thông tin hữu ích','Lưu ý tra cứu'],[
('Database','Database → Database Instances. Đọc ID/Name, Datastore, Datastore Version, Host, Database Disk và Status.','Phiên khảo sát có DB1, MySQL 8.0.46, Active. Không dùng Delete hoặc More để thử xử lý.'),
('Bản sao database','Database → Backups; Configuration Groups nằm trong cùng menu.','Chọn đúng loại bản sao database, phân biệt với Volume Backups.'),
('Kubernetes','Container → K8s Clusters hoặc K8s Templates. Bảng cluster có Status, Health Status và Keypair.','Trang K8s Clusters báo Get clusters error. Không diễn giải No Data khi có lỗi là không có cluster.'),
('Orchestration','Orchestration → Stacks, Resource Types, Template Versions, Template Generator.','Danh mục hiện hữu; các luồng triển khai stack không nằm trong quy trình này.'),
('Image và cấu hình máy','Compute → Images, Flavors, Key Pairs, Server Groups, Instance Snapshots.','Dùng để tra cứu cấu hình tham chiếu của máy ảo. Không tạo key hoặc tải khóa riêng để kiểm tra.'),
('Tự co giãn','Compute → Autoscale Rules (V), Autoscale History, Auto Scaling Groups, ASG History.','Ưu tiên xem lịch sử. Thay đổi rule hoặc group có thể ảnh hưởng số lượng máy chạy.'),
('Lịch sao lưu','Storage → Snapshot Schedules và Backup Schedules.','Xem thông tin lịch hiện có; tạo hoặc sửa lịch sẽ làm thay đổi vận hành.'),
('Thùng rác','Maintenance → Recycle Bin.','Không thực hiện khôi phục hoặc xóa vĩnh viễn trong bước tra cứu.')],[1.35,3.3,2.45])
h('Đối chiếu trên Horizon')
p('Project có các nhóm Databases, Container Infra và Orchestration. Chỉ mở các mục được quyền xem. Tên menu hoặc việc có endpoint trong System Information không đủ xác nhận dịch vụ đã sẵn sàng cho mọi thao tác.')
p('Nếu trang báo lỗi, ghi lại portal, project, đường dẫn, thông báo nguyên văn và thời điểm. Dùng màn hình tương ứng đã hoạt động hoặc nhờ vận hành kiểm tra dịch vụ; không tự tạo tài nguyên thử để xác minh.')

page(15,'Horizon dành cho nhân sự vận hành')
p('Nhóm Admin dành cho tài khoản được cấp quyền quản trị. Không dùng số liệu tại đây để thay thế quota của một project: phạm vi tài nguyên có thể bao gồm nhiều project.')
h('Hypervisor và tài nguyên tổng hợp')
steps(['Mở Admin → Compute → Hypervisors.', 'Đọc Hypervisor Summary và Resource Providers Summary để xem RAM, local disk và vCPU theo host.', 'Trong bảng Hypervisor, đối chiếu Hostname, Type, RAM, Local Storage và Instances. Dùng tên host trong chi tiết máy ảo để tìm đúng dòng.'])
p('Phiên khảo sát hiển thị 3 hypervisor QEMU: compute-01, compute-02 và compute-03. Bảng ghi lần lượt 15, 7 và 3 instance trên các host. Các giá trị tổng hợp này khác phạm vi 5 máy ảo của project admin.')
h('Trạng thái các dịch vụ')
steps(['Mở Admin → System → System Information.', 'Tab Services liệt kê tên dịch vụ, loại, region và endpoint. Chỉ có endpoint chưa chứng minh dịch vụ khỏe.', 'Mở Compute Services, Block Storage Services hoặc Network Agents. Đọc Status, State và Last Updated cùng nhau.'])
picture('35-horizon-compute-services','Compute Services trên Horizon với Status State và Last Updated')
p('Khi khảo sát, 9 Compute Services, 9 Block Storage Services và 15 Network Agents đều hiển thị Enabled và Up. Đây là trạng thái do portal báo tại thời điểm xem, không phải kiểm thử xử lý yêu cầu đầu cuối. Horizon hiển thị phiên bản 25.7.4.')

page(16,'Đối chiếu đường dẫn giữa hai portal')
p('Hai portal cùng hiển thị tài nguyên cloud trong project đang chọn. Dùng ID tài nguyên để đối chiếu. Tên và thứ tự cột có thể khác nhau; không mặc định mọi chức năng có bản tương ứng trên portal còn lại.')
table(['Nhu cầu','XOR','Horizon trong nhóm Project'],[
('Quota project','Home','Compute → Overview → Limit Summary'),
('Máy ảo','Compute → Instances','Compute → Instances'),
('Volume','Storage → Volumes','Volumes → Volumes'),
('Snapshot volume','Storage → Volume Snapshots','Volumes → Snapshots'),
('Backup volume','Storage → Volume Backups','Volumes → Backups'),
('Mạng','Network → Networks','Network → Networks'),
('Router','Network → Routers','Network → Routers'),
('Sơ đồ mạng','Network → Topology','Network → Network Topology'),
('Floating IP','Network → Floating IPs','Network → Floating IPs'),
('Security Group','Network → Security Groups','Network → Security Groups'),
('Load Balancer','Network → Load Balancers','Network → Load Balancers'),
('VPN','Network → VPNs','Network → VPN'),
('Database','Database','Databases'),
('Kubernetes','Container','Container Infra'),
('Stack','Orchestration','Orchestration')],[1.65,2.55,2.9])
h('Chức năng chuyên biệt')
p('Security Center, X-OR SECURE và Monitor Center đã được quan sát trên XOR. Hypervisors và các bảng trạng thái dịch vụ đã được quan sát trong nhóm Admin của Horizon. Không có đường dẫn đối ứng đã xác nhận cho các nhóm này trong khảo sát.')
p('Bảng là bản đồ menu. Các mục chỉ quan sát ở cấp menu gồm backup, snapshot, load balancer, VPN và orchestration; nội dung chi tiết và luồng xử lý của các mục đó chưa được kiểm chứng trong phiên này.')

page(17,'Xử lý vướng mắc khi tra cứu')
table(['Hiện tượng','Việc người dùng nên kiểm tra','Thông tin cần gửi vận hành'],[
('Không thấy tài nguyên hoặc menu','Đúng portal, project, domain và tiêu chí lọc; chờ trang tải xong. Quyền tài khoản có thể giới hạn dữ liệu.','Tên/ID cần tìm, project, menu, thời điểm và ảnh vùng báo lỗi.'),
('No Data đi kèm lỗi','Đọc thông báo lỗi trước; không kết luận tài nguyên bị xóa.','Nguyên văn lỗi và đường dẫn. Ví dụ khảo sát: K8s Clusters báo Get clusters error.'),
('Số đếm giữa các trang khác nhau','Đối chiếu phạm vi project/toàn cloud, cửa sổ thời gian, dữ liệu lịch sử và thời điểm cập nhật.','Tên hai trang, bộ lọc, số đếm và thời gian chụp.'),
('Máy ảo Active nhưng ứng dụng lỗi','Ghi IP, trạng thái cloud, volume và Security Group; chuyển người quản trị ứng dụng kiểm tra dịch vụ.','ID máy, ứng dụng/cổng cần truy cập, nguồn truy cập và biểu hiện lỗi.'),
('Cảnh báo bằng 0 hoặc số liệu cũ','Đọc Health, connector, heartbeat, lần scan và phạm vi dữ liệu.','Dịch vụ, trạng thái nguồn, lần cập nhật cuối và tài sản bị ảnh hưởng.'),
('Quota gần đầy','Đọc Used/Limit cho đúng loại tài nguyên. Không xóa tài nguyên để tự giải phóng quota.','Project, quota đang dùng/giới hạn và nhu cầu bổ sung.'),
('Phiên JumpServer gián đoạn','Ghi tên tab/tài sản và thông báo phiên; liên hệ quản trị truy cập nếu không vào lại được bằng quy trình đã cấp.','Thời điểm, tên phiên và lỗi. Không gửi mật khẩu hoặc token.')],[1.5,3.25,2.35])
h('Những điểm cần đối chiếu sau khảo sát')
p('Monitor Overview hiển thị một số bộ đếm mạng bằng 0 trong khi danh sách tài nguyên của project có dữ liệu. K8s Clusters báo lỗi dù Home có chỉ số cluster. Hai điểm này cần được vận hành xác minh nguồn dữ liệu và dịch vụ; khảo sát chỉ xem chưa xác định nguyên nhân.')
p('Security Center có các dịch vụ chưa kết nối hoặc chưa hoạt động như mô tả ở phần 12. Riêng SIEM hiển thị Storage Not Measured và Retention Not Configured; cần kiểm tra cấu hình tại nguồn trước khi đưa ra kết luận về dung lượng hoặc thời gian lưu log.')

page(18,'Quy trình kiểm tra và gửi yêu cầu hỗ trợ')
h('Trình tự kiểm tra chỉ đọc')
steps(['Xác nhận tài khoản, project và portal; ghi thời gian bắt đầu.', 'Đọc quota và danh sách máy ảo. Ghi trạng thái bất thường, ID và IP liên quan.', 'Kiểm tra volume gắn với máy, mạng, router, Floating IP và Security Group khi có vấn đề kết nối hoặc lưu trữ.', 'Xem Action Logs hoặc Audit Logs để tìm mốc thời gian và Request ID có liên quan.', 'Trong Security Center, kiểm tra nguồn dữ liệu và thời điểm cập nhật trước khi đọc phát hiện.', 'Ghi nhận kết quả, chuyển yêu cầu cho người phụ trách. Mọi thay đổi được thực hiện theo quy trình vận hành riêng.'])
h('Mẫu thông tin hỗ trợ')
p('Tiêu đề: [Portal] [Project] [Tài nguyên] – biểu hiện cần hỗ trợ')
p('Thời điểm và múi giờ: …\nNgười yêu cầu và kênh liên hệ nội bộ: …\nPortal, project và đường dẫn menu: …\nTên tài nguyên, ID và IP liên quan: …\nTrạng thái đang thấy và thông báo lỗi nguyên văn: …\nThời điểm cập nhật cuối hoặc Request ID nếu có: …\nẢnh màn hình đã loại bỏ thông tin bí mật: …\nNhu cầu mong muốn và ảnh hưởng nghiệp vụ: …\nCác bước tra cứu đã thực hiện: …')
h('Phạm vi kiểm chứng của tài liệu')
p('Nguồn là giao diện thực tế của hai phiên WEB-XOR-Portal và WEB-Horizon trong JumpServer, khảo sát ngày 13/10/2026. Đã xem danh sách và một số chi tiết máy ảo, volume, mạng, quy tắc bảo mật, Floating IP, quota, monitoring, audit, các trang tổng quan Security Center, Posture Findings, SIEM Alerts, trang thiết kế, trang TCVN Cấp độ 3, database, Kubernetes và các trang hạ tầng Horizon.')
p('Các tab hoặc menu còn lại được mô tả để định vị chức năng, không khẳng định đã kiểm chứng toàn bộ nội dung. Chưa kiểm thử tạo/sửa/xóa tài nguyên, phục hồi bản sao, đăng nhập console, khả năng phục vụ ứng dụng hay các luồng triển khai và ứng cứu. Tài liệu không chứa mật khẩu, token hoặc khóa riêng.')
p('Ảnh minh họa có tên tài nguyên, địa chỉ mạng và thông tin vận hành nội bộ. Chỉ lưu hành trong phạm vi được phép; thay ảnh bằng dữ liệu đã được phép công bố trước khi cung cấp ra ngoài.')

doc.save(OUT)
print(OUT)
